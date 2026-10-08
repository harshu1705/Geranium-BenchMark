# -*- coding: utf-8 -*-
from docx import Document

doc = Document()

def h(text, level=1):
    doc.add_heading(text, level=level)

def p(text):
    doc.add_paragraph(text)

def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    for i, htext in enumerate(headers):
        t.rows[0].cells[i].text = htext
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v)
    doc.add_paragraph()

doc.add_heading('Tri-County Community Bank — Q4 2025 ALLL Reserve Adequacy Memorandum', level=0)
p('Prepared by: Credit Risk Division, Tri-County Community Bank')
p(
    "Review Date: January 12, 2026 — the date Exception EXC-001's covenant waiver was issued, the most "
    "recent dated entry on file in Exception_Tracking_Log_Q4_2025.xlsx at the time this testing was performed."
)

h('1. Executive Summary')
p(
    "Auditors are going to want to see this before quarter close, so here is where the reserve actually "
    "stands. The commercial book runs 29 loans, $134,656,509 outstanding in total, and under the ALLL/CECL "
    "Reserve Policy the required reserve against that book comes to $1,651,606 — $1,310,966 pooled across "
    "Risk Rating tiers 1 through 5, plus a $340,640 specific reserve carried against LN-008 (Palmetto Print "
    "& Packaging Inc.), which is individually evaluated because it is collateral-dependent following its "
    "Chapter 11 filing."
)
p(
    "Booked ALLL sits at $1,250,000. That leaves a shortfall of $401,606, or about 32% of what's currently "
    "on the books — well outside the 10% band Policy Section 6 allows. An additional provision of $401,606 "
    "should go in this quarter, and the Board Audit Committee needs to see it."
)

h('2. Portfolio Segmentation')
p(
    "Every loan was bucketed by Internal Risk Rating straight from Loan_Register_Q4_2025.xlsx. Ratings 1 "
    "through 5 get pooled treatment; Rating 6 doesn't — Policy Section 4 pulls it out for individual "
    "evaluation, so it's excluded from the pool table below."
)
p(
    "The six segments were re-tallied from scratch rather than pulled from any existing summary, and they "
    "tie out exactly: $134,656,509 across all 29 loans, no omissions, nothing double-counted."
)
table(
    ['Risk Rating', 'Loan Count', 'Aggregate Balance', 'Treatment'],
    [
        [1, 1, '$21,117,450', 'Pooled'],
        [2, 19, '$87,998,690', 'Pooled'],
        [3, 6, '$13,420,408', 'Pooled'],
        [4, 1, '$3,817,432', 'Pooled'],
        [5, 1, '$7,214,889', 'Pooled'],
        [6, 1, '$1,087,640', 'Individually evaluated (LN-008)'],
    ]
)

h('3. Pooled Reserve Calculation')
p(
    "Base loss rates came from ALLL_Loss_Rate_and_Reserve_Data.xlsx. Before applying any of them, each pool "
    "got checked against the Covenant Compliance Workbook, the Borrower Financial Statements, and the "
    "Appraisal and Collateral Summary — Policy Section 5 requires an active, un-waived breach before a "
    "25bps Q-Factor gets tacked on, and pre-calculated breach flags aren't good enough on their own. DSCR "
    "was rebuilt as Net Operating Income over Total Annual Debt Service, LTV as Outstanding Balance over "
    "Appraised Value, and Current Ratio as Current Assets over Current Liabilities, all pulled from the "
    "underlying files rather than taken on faith from the workbook."
)
p(
    "Three loans in Pool 3 are breaching with nothing on file to offset them: LN-004 (Coastal Auto Group, "
    "DSCR 1.14x against a 1.20x floor), LN-006 (Harbor Grille Restaurant Group, Current Ratio 1.03x against "
    "1.20x), and LN-010 (U-Save Self Storage LLC, LTV 85.2% against a 75% cap). None has a waiver, so the "
    "Q-Factor applies to the whole pool. Pool 5 has one breaching loan in the same boat — LN-002 (Corrigan "
    "Retail Partners), LTV 87.9% against 75% and Current Ratio 0.94x against 1.00x — no waiver there either, "
    "so Pool 5 gets the add-on too."
)
p(
    "Pool 4 is different. Its only loan, LN-001 (Talbot Steel Fabricators LLC), is breaching DSCR at 1.07x "
    "versus the 1.20x minimum, but Exception EXC-001 shows an approved waiver issued January 12, 2026 and "
    "running through March 31, 2026 — active and not expired as of this review. Policy Section 5 is explicit "
    "that a breach \"offset by a documented, approved waiver\" doesn't trigger the Q-Factor, so Pool 4 stays "
    "at its base rate. Pools 1 and 2 aren't breaching anything and don't get an add-on either."
)
p('Pooled reserve across the five tiers comes to $1,310,966.')
table(
    ['Risk Rating', 'Pool Balance', 'Base Rate', 'Q-Factor Trigger', 'Adjusted Rate', 'Pooled Reserve'],
    [
        [1, '$21,117,450', '0.10%', 'None', '0.10%', '$21,117'],
        [2, '$87,998,690', '0.35%', 'None', '0.35%', '$307,995'],
        [3, '$13,420,408', '0.90%', 'LN-004, LN-006, LN-010 breaching, no waiver', '1.15%', '$154,335'],
        [4, '$3,817,432', '3.25%', 'LN-001 breaching but waived (EXC-001, exp. 3/31/2026)', '3.25%', '$124,067'],
        [5, '$7,214,889', '9.50%', 'LN-002 breaching, no waiver', '9.75%', '$703,452'],
    ]
)

h('4. Individually Evaluated Loan Analysis')
p(
    "LN-008 (Palmetto Print & Packaging Inc.) is Rating 6 and 97 days past due — either one alone would pull "
    "it out of the pool under Policy Section 4, and it hits both. The borrower filed Chapter 11 on December "
    "9, 2025, and at this point repayment is expected to come from liquidating the manufacturing equipment "
    "pledged as collateral rather than from operating cash flow."
)
p(
    "Exception EXC-010 backs this up independently, noting an active loss reserve on this loan, and the "
    "Appraisal and Collateral Summary flags the Chapter 11 filing in its notes as well. No other loan in the "
    "book is Rating 6 or 90-plus days past due, so LN-008 is the only one that needed this treatment."
)
table(
    ['Line Item', 'Amount'],
    [
        ['Outstanding Balance', '$1,087,640'],
        ['Appraised Collateral Value (Dec. 2025 inspection)', '$830,000'],
        ['Estimated Disposition Cost (10%)', '$83,000'],
        ['Net Realizable Collateral Value', '$747,000'],
        ['Specific Reserve (Balance minus Net Realizable Value)', '$340,640'],
    ]
)

h('5. Covenant Breach and Watch List Cross-Reference')
p(
    "Five loans are carrying an active breach right now: LN-001, LN-002, LN-004, LN-006, and LN-010. LN-008 "
    "isn't part of this list because it's individually evaluated in Section 4, not pooled. Each one's "
    "Exception Tracking Log entry was pulled and checked for a Waiver Issued flag and expiry date before any "
    "Q-Factor got applied. Only LN-001 has an approved waiver still in force; the other four all show "
    "\"Waiver Issued: No\" and stay subject to the add-on."
)
p(
    "LN-001, LN-002, and LN-008 also carry a Watch List designation, which under the Loan Portfolio Review "
    "Policy just means an Internal Risk Rating of 4 or higher — that's a separate trigger from the Q-Factor "
    "and doesn't change the pooled math on its own, since the base loss rate already prices in the elevated "
    "rating. LN-001 stays on the Watch List even with the waiver in place; the waiver only affects the "
    "Q-Factor treatment, not the rating-based classification. LN-004, LN-006, and LN-010 are Rating 3, below "
    "the Watch List cutoff, but their breaches still pull the Pool 3 Q-Factor because that add-on applies to "
    "the whole pool, not just the specific loan that's breaching."
)

h('6. Reserve Adequacy Conclusion')
p(
    "$401,606 works out to about 32% of the currently booked reserve, well past the 10% band Policy Section "
    "6 allows either direction. That makes the reserve inadequate as booked — an additional provision of "
    "$401,606 needs to go in this quarter and get reported up to the Board Audit Committee. Most of the gap "
    "traces back to two credits: LN-002, where a single loan in the Rating 5 pool is driving $703,452 of "
    "required reserve on its own, and LN-008's $340,640 specific reserve. Both reflect real, recent "
    "deterioration — a payment demand and two concurrent covenant breaches on LN-002, and a bankruptcy filing "
    "with falling collateral value on LN-008 — that the current booked balance doesn't seem to have caught "
    "up with yet."
)

h('7. Policy/Source Reconciliation')
p(
    "The loss rate table, the disposition cost assumption, the Q-Factor rate, and the booked ALLL figure all "
    "came straight from ALLL_Loss_Rate_and_Reserve_Data.xlsx and didn't conflict with anything else in the "
    "file set. Pool composition was rebuilt from Loan_Register_Q4_2025.xlsx rather than taken from any prior "
    "summary, and the six segments tie back to the full 29-loan, $134,656,509 portfolio."
)
p(
    "None of the breach determinations feeding the Q-Factor decision were taken on faith from the Covenant "
    "Compliance Workbook's actual-vs-threshold columns. DSCR came from Net Operating Income divided by Total "
    "Annual Debt Service in Borrower_Financial_Statements_Summary.xlsx; LTV came from the Loan Register "
    "balance divided by the appraised value in the Appraisal and Collateral Summary; Current Ratio came from "
    "Current Assets over Current Liabilities in the financial statements file. Everything recomputed this way "
    "landed within normal rounding of the workbook's own figures — no real conflicts turned up."
)
p(
    "Breach status by itself wasn't treated as enough to justify a Q-Factor, though. Every breaching loan's "
    "Waiver Issued and Waiver Expiry Date fields in Exception_Tracking_Log_Q4_2025.xlsx got checked first, "
    "consistent with Policy Section 5's carve-out for breaches \"offset by a documented, approved waiver.\" "
    "That check is what surfaced LN-001's waiver — EXC-001, issued January 12, 2026, running to March 31, "
    "2026, still active at the time of this review — so Pool 4 was reserved at base rate with no add-on. "
    "LN-002, LN-004, LN-006, and LN-010 came back with \"Waiver Issued: No\" on their respective exceptions "
    "and were assessed the Q-Factor accordingly."
)
p(
    "For LN-008's specific reserve, the appraised value used was the most recent one on file — the December "
    "2025 inspection in the Appraisal and Collateral Summary, $830,000 — rather than anything from "
    "origination, which is what Policy Section 4 calls for regardless of how old the loan is. Exception "
    "EXC-010 independently backs this up, noting the loss reserve as already active for this credit."
)

h('8. Approval and Distribution')
p(
    "Per Policy Section 8, this memorandum must be reviewed and approved by the Chief Credit Officer prior "
    "to submission to the Board Audit Committee. It is submitted below for that review."
)
table(
    ['Role', 'Name', 'Signature', 'Date'],
    [
        ['Prepared by', 'Credit Risk Division', '', 'January 12, 2026'],
        ['Reviewed and Approved by', 'Chief Credit Officer', '', ''],
    ]
)
p(
    "Distribution per Policy Section 8: Chief Credit Officer, Chief Financial Officer, Board Audit Committee, "
    "and the Bank's external auditor."
)

doc.save('check_golden/Tri_County_ALLL_Reserve_Adequacy_Memo_Q4_2025.docx')
print('saved')
