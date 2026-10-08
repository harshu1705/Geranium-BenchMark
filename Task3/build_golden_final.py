# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt

doc = Document()

def h(text, level=1):
    doc.add_heading(text, level=level)

def p(text):
    doc.add_paragraph(text)

def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    for i, htext in enumerate(headers):
        t.rows[0].cells[i].text = htext
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v)
    doc.add_paragraph()

title = doc.add_heading('Tri-County Community Bank — Q4 2025 ALLL Reserve Adequacy Memorandum', level=0)
p('Prepared by: Credit Risk Division, Tri-County Community Bank')
p('Review Date: January 23, 2026 (the most recent loan review date recorded in Loan_Register_Q4_2025.xlsx, LN-028)')

h('1. Executive Summary')
p(
    "This memorandum tests the adequacy of Tri-County Community Bank's Allowance for Loan and Lease Losses "
    "(ALLL) for the Q4 2025 commercial loan portfolio under the methodology set out in the ALLL/CECL Reserve "
    "Policy. The portfolio comprises 29 active commercial loans with an aggregate outstanding balance of "
    "$134,656,509."
)
p(
    "The total reserve required under this methodology is $1,651,606, combining $1,310,966 in pooled reserves "
    "across Risk Rating pools 1 through 5 and a $340,640 specific reserve for the one individually evaluated "
    "loan, LN-008 (Palmetto Print & Packaging Inc.), which is collateral-dependent following its Chapter 11 "
    "filing. Against a currently booked ALLL balance of $1,250,000, the portfolio carries a reserve shortfall "
    "of $401,606, or 32.1% of the booked balance. This shortfall exceeds the 10% adequacy tolerance band in "
    "Policy Section 6, and an additional provision of $401,606 should be recognized in the current quarter."
)

h('2. Portfolio Segmentation')
p(
    "All 29 loans were segmented by Internal Risk Rating as recorded in Loan_Register_Q4_2025.xlsx. Ratings 1 "
    "through 5 are pooled; Rating 6 is individually evaluated per Policy Section 4 and excluded from the pooled "
    "segmentation below."
)
p(
    "Segment counts and balances were independently re-tallied from the Loan Register rather than taken from "
    "any pre-existing summary; the six segments sum to the full portfolio balance of $134,656,509 and the full "
    "loan count of 29, confirming no loan was omitted or double-counted."
)
table(
    ['Risk Rating', 'Loan Count', 'Aggregate Balance', 'Treatment'],
    [
        [1, 1, '$21,117,450', 'Pooled'],
        [2, 19, '$87,998,690', 'Pooled'],
        [3, 6, '$13,420,408', 'Pooled'],
        [4, 1, '$3,817,432', 'Pooled'],
        [5, 1, '$7,214,889', 'Pooled'],
        [6, 1, '$1,087,640', 'Individually evaluated (LN-008)'],
    ]
)

h('3. Pooled Reserve Calculation')
p(
    "Base historical loss rates were taken from ALLL_Loss_Rate_and_Reserve_Data.xlsx. Before applying them, "
    "each pool was checked against the Covenant Compliance Workbook, Borrower Financial Statements, and "
    "Appraisal and Collateral Summary to determine whether it contains an active, un-waived covenant breach "
    "requiring a Policy Section 5 Q-Factor add-on of 25 basis points. Breach status was independently "
    "recomputed rather than accepted from the workbook's pre-calculated fields: DSCR was recomputed as Net "
    "Operating Income divided by Total Annual Debt Service, LTV as Outstanding Balance divided by Appraised "
    "Value, and Current Ratio as Current Assets divided by Current Liabilities."
)
p(
    "Pool 3 contains three breaching loans with no approved waiver on file in the Exception Tracking Log: "
    "LN-004 (Coastal Auto Group, DSCR 1.14x vs. 1.20x), LN-006 (Harbor Grille Restaurant Group, Current Ratio "
    "1.03x vs. 1.20x), and LN-010 (Wexford Storage Partners, LTV 85.2% vs. 75%). All three receive the "
    "Q-Factor. Pool 5 contains one breaching loan with no approved waiver, LN-002 (Corrigan Retail Partners, "
    "LTV 87.9% vs. 75% and Current Ratio 0.94x vs. 1.00x), and also receives the Q-Factor."
)
p(
    "Pool 4's sole loan, LN-001 (Talbot Steel Fabricators LLC), carries an active DSCR breach (1.07x vs. "
    "1.20x), but Exception EXC-001 in the Exception Tracking Log records an approved waiver issued January 12, "
    "2026, expiring March 31, 2026 — active and unexpired as of this review date. Policy Section 5 exempts "
    "any breach \"offset by a documented, approved waiver\" from the Q-Factor trigger. Pool 4 therefore "
    "receives no Q-Factor add-on and is reserved at its base historical loss rate only. Pools 1 and 2 contain "
    "no loans with an active covenant breach and likewise receive no Q-Factor add-on."
)
p('Total pooled reserve across all five pools: $1,310,966.')
table(
    ['Risk Rating', 'Pool Balance', 'Base Rate', 'Q-Factor Trigger', 'Adjusted Rate', 'Pooled Reserve'],
    [
        [1, '$21,117,450', '0.10%', 'None', '0.10%', '$21,117'],
        [2, '$87,998,690', '0.35%', 'None', '0.35%', '$307,995'],
        [3, '$13,420,408', '0.90%', 'LN-004, LN-006, LN-010 breaching, no waiver', '1.15%', '$154,335'],
        [4, '$3,817,432', '3.25%', 'LN-001 breaching but waived (EXC-001, exp. 3/31/2026)', '3.25%', '$124,067'],
        [5, '$7,214,889', '9.50%', 'LN-002 breaching, no waiver', '9.75%', '$703,452'],
    ]
)

h('4. Individually Evaluated Loan Analysis')
p(
    "LN-008 (Palmetto Print & Packaging Inc.) carries an Internal Risk Rating of 6 and is 97 days past due, "
    "meeting both individual-evaluation triggers under Policy Section 4. The loan is collateral-dependent: the "
    "borrower filed Chapter 11 on December 9, 2025, and repayment is expected to be provided through "
    "liquidation of the manufacturing equipment pledged as collateral."
)
p(
    "This specific reserve requirement is corroborated by Exception EXC-010 in the Exception Tracking Log, "
    "which independently records a loss reserve as active for this loan, and by the Appraisal and Collateral "
    "Summary note referencing the Chapter 11 filing. No other loan in the portfolio is rated 6 or 90+ days past "
    "due; no other loan requires individual evaluation."
)
table(
    ['Line Item', 'Amount'],
    [
        ['Outstanding Balance', '$1,087,640'],
        ['Appraised Collateral Value (Dec. 2025 inspection)', '$830,000'],
        ['Estimated Disposition Cost (10%)', '$83,000'],
        ['Net Realizable Collateral Value', '$747,000'],
        ['Specific Reserve (Balance minus Net Realizable Value)', '$340,640'],
    ]
)

h('5. Covenant Breach and Watch List Cross-Reference')
p(
    "Five loans carry an active covenant breach as of the review date: LN-001, LN-002, LN-004, LN-006, and "
    "LN-010 (LN-008 is excluded from this pool-level analysis because it is individually evaluated in Section "
    "4, not pooled). Every breaching loan's Exception Tracking Log record was checked for a Waiver Issued flag "
    "and Waiver Expiry Date before the Q-Factor was applied. Only LN-001 has an approved, unexpired waiver "
    "(EXC-001); LN-002, LN-004, LN-006, and LN-010 show \"Waiver Issued: No\" in the Exception Tracking Log and "
    "therefore remain subject to the Q-Factor."
)
p(
    "LN-001, LN-002, and LN-008 additionally carry a Watch List designation under the Loan Portfolio Review "
    "Policy (Internal Risk Rating of 4 or higher). The Watch List designation itself does not directly change "
    "the pooled reserve calculation — the Risk Rating pool already reflects elevated risk through its base "
    "historical loss rate — but the presence of an un-waived breach within a pool is what triggers the "
    "Q-Factor add-on documented in Section 3. LN-001 remains on the Watch List by risk rating even though its "
    "breach is currently waived; the waiver affects only the Q-Factor treatment, not the Watch List "
    "classification. LN-004, LN-006, and LN-010 are not on the Watch List (their Internal Risk Ratings are 3, "
    "below the Rating ≥ 4 threshold), but their un-waived breaches still trigger the Pool 3 Q-Factor because "
    "the Q-Factor is a pool-level, not loan-level, adjustment."
)

h('6. Reserve Adequacy Conclusion')
p(
    "The $401,606 shortfall represents 32.1% of the currently booked balance, well outside the ±10% "
    "adequacy tolerance band in Policy Section 6. The reserve is therefore inadequate, and an additional "
    "provision of $401,606 should be recognized in the current quarter and reported to the Board Audit "
    "Committee. The shortfall is driven primarily by the Rating 5 pool (LN-002, $703,452 of required reserve "
    "on a single loan) and the LN-008 specific reserve ($340,640), both of which reflect deterioration — a "
    "payment demand and dual covenant breach on LN-002, and a Chapter 11 filing with declining collateral "
    "value on LN-008 — that the currently booked balance does not appear to fully capture."
)

h('7. Policy/Source Reconciliation')
p(
    "Historical loss rates, the disposition cost assumption, the Q-Factor rate, and the currently booked ALLL "
    "balance were taken directly from ALLL_Loss_Rate_and_Reserve_Data.xlsx as pre-calculated policy inputs; no "
    "source conflicts were found among these figures. Risk Rating pool composition was independently re-tallied "
    "from Loan_Register_Q4_2025.xlsx rather than assumed, and the six segments were confirmed to sum to the "
    "full 29-loan, $134,656,509 portfolio."
)
p(
    "Covenant breach status feeding the Q-Factor determination was independently recomputed rather than read "
    "from the Covenant Compliance Workbook's pre-calculated actual-vs-threshold fields: DSCR was recomputed "
    "from Borrower_Financial_Statements_Summary.xlsx (Net Operating Income ÷ Total Annual Debt Service), "
    "LTV from the Loan Register outstanding balance divided by the Appraisal and Collateral Summary appraised "
    "value, and Current Ratio from Current Assets ÷ Current Liabilities in the financial statements file. "
    "All recomputed values matched the workbook's pre-calculated figures within normal rounding tolerance "
    "(±0.01x or ±0.1 percentage point), and no conflicts were identified between the pre-calculated "
    "workbook values and the independently derived calculations."
)
p(
    "Critically, breach status alone was not treated as sufficient grounds for a Q-Factor: every breaching "
    "loan's Waiver Issued and Waiver Expiry Date fields in Exception_Tracking_Log_Q4_2025.xlsx were reviewed "
    "before applying the add-on, per Policy Section 5's exemption for breaches \"offset by a documented, "
    "approved waiver.\" This review identified that LN-001's DSCR breach (EXC-001) carries an approved waiver "
    "issued January 12, 2026 and expiring March 31, 2026, which remains active and unexpired as of this "
    "January 23, 2026 review date; Pool 4 was therefore reserved at its base rate with no Q-Factor add-on. "
    "LN-002, LN-004, LN-006, and LN-010 show no waiver on file (Waiver Issued: No) and their pools were "
    "assessed the Q-Factor accordingly."
)
p(
    "The LN-008 specific reserve calculation used the most recent appraised collateral value on file (the "
    "December 2025 collateral inspection in the Appraisal and Collateral Summary, $830,000) rather than an "
    "older or origination-date value, consistent with Policy Section 4's requirement to use the most recent "
    "appraisal regardless of loan origination date. This treatment is corroborated by Exception EXC-010, which "
    "independently documents an active loss reserve for this loan in the Exception Tracking Log."
)

h('8. Approval and Distribution')
p(
    "This memorandum has been prepared in accordance with Policy Section 8 and is submitted for review and "
    "approval prior to distribution."
)
table(
    ['Role', 'Name', 'Signature', 'Date'],
    [
        ['Prepared by', 'Credit Risk Division', '', 'January 23, 2026'],
        ['Reviewed and Approved by', 'Chief Credit Officer', '', ''],
    ]
)
p('Distribution: Chief Credit Officer, Chief Financial Officer, Board Audit Committee, External Auditor.')

doc.save('check_golden/Tri_County_ALLL_Reserve_Adequacy_Memo_Q4_2025.docx')
print('saved')
