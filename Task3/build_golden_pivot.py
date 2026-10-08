import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
import os

BASE = r'c:\Users\user\OneDrive\Desktop\geranium project\Task3'
OUTDIR = os.path.join(BASE, 'pivot_golden')
os.makedirs(OUTDIR, exist_ok=True)

d = Document()
d.add_heading('Tri-County Community Bank — Q4 2025 ALLL Reserve Adequacy Memorandum', level=0)
d.add_paragraph('Prepared by Credit Risk Division | Review Date: January 23, 2026')

d.add_heading('1. Executive Summary', level=1)
d.add_paragraph(
    "This memorandum tests the adequacy of Tri-County Community Bank's Allowance for Loan and Lease "
    "Losses (ALLL) for the Q4 2025 commercial loan portfolio under the methodology set out in the ALLL/"
    "CECL Reserve Policy. The portfolio comprises 29 active commercial loans with an aggregate outstanding "
    "balance of $134,656,509."
)
d.add_paragraph(
    "The total reserve required under this methodology is $1,661,149, combining $1,320,509 in pooled "
    "reserves across Risk Rating pools 1 through 5 and a $340,640 specific reserve for the one "
    "individually evaluated loan, LN-008 (Palmetto Print & Packaging Inc.), which is collateral-dependent "
    "following its Chapter 11 filing. Against a currently booked ALLL balance of $1,250,000, the portfolio "
    "carries a reserve shortfall of $411,149. This shortfall exceeds the 10% adequacy tolerance band in "
    "Policy Section 6, and an additional provision of $411,149 should be recognized in the current quarter."
)

d.add_heading('2. Portfolio Segmentation', level=1)
d.add_paragraph(
    "All 29 loans were segmented by Internal Risk Rating as recorded in Loan_Register_Q4_2025.xlsx. Ratings "
    "1 through 5 are pooled; Rating 6 is individually evaluated per Policy Section 4 and excluded from the "
    "pooled segmentation below."
)
table = d.add_table(rows=1, cols=4)
table.style = 'Light Grid Accent 1'
hdr = table.rows[0].cells
hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = 'Risk Rating', 'Loan Count', 'Aggregate Balance', 'Treatment'
seg_rows = [
    ('1', '1', '$21,117,450', 'Pooled'),
    ('2', '19', '$87,998,690', 'Pooled'),
    ('3', '6', '$13,420,408', 'Pooled'),
    ('4', '1', '$3,817,432', 'Pooled'),
    ('5', '1', '$7,214,889', 'Pooled'),
    ('6', '1', '$1,087,640', 'Individually evaluated (LN-008)'),
]
for row in seg_rows:
    cells = table.add_row().cells
    for i, v in enumerate(row):
        cells[i].text = v
d.add_paragraph(
    "Segment counts and balances were independently re-tallied from the Loan Register rather than taken "
    "from any pre-existing summary; the six segments sum to the full portfolio balance of $134,656,509 "
    "and the full loan count of 29, confirming no loan was omitted or double-counted."
)

d.add_heading('3. Pooled Reserve Calculation', level=1)
d.add_paragraph(
    "Base historical loss rates were taken from ALLL_Loss_Rate_and_Reserve_Data.xlsx. Before applying "
    "them, each pool was checked against the Covenant Compliance Workbook, Borrower Financial Statements, "
    "and Appraisal and Collateral Summary to determine whether it contains an active covenant breach "
    "requiring a Policy Section 5 Q-Factor add-on of 25 basis points. Breach status was independently "
    "recomputed rather than accepted from the workbook's pre-calculated fields: DSCR was recomputed as Net "
    "Operating Income divided by Total Annual Debt Service, LTV as Outstanding Balance divided by Appraised "
    "Value, and Current Ratio as Current Assets divided by Current Liabilities. Pool 3 contains three "
    "breaching loans (LN-004 DSCR 1.14x vs. 1.20x, LN-006 Current Ratio 1.03x vs. 1.20x, LN-010 LTV 85.2% "
    "vs. 75%); Pool 4 contains one breaching loan (LN-001, DSCR 1.07x vs. 1.20x); Pool 5 contains one "
    "breaching loan (LN-002, LTV 87.9% vs. 75% and Current Ratio 0.94x vs. 1.00x). Pools 1 and 2 contain no "
    "loans with an active covenant breach and receive no Q-Factor add-on."
)
table2 = d.add_table(rows=1, cols=6)
table2.style = 'Light Grid Accent 1'
hdr2 = table2.rows[0].cells
for i,h in enumerate(['Risk Rating', 'Pool Balance', 'Base Rate', 'Q-Factor Trigger', 'Adjusted Rate', 'Pooled Reserve']):
    hdr2[i].text = h
pool_rows = [
    ('1', '$21,117,450', '0.10%', 'None', '0.10%', '$21,117'),
    ('2', '$87,998,690', '0.35%', 'None', '0.35%', '$307,995'),
    ('3', '$13,420,408', '0.90%', 'LN-004, LN-006, LN-010 breaching', '1.15%', '$154,335'),
    ('4', '$3,817,432', '3.25%', 'LN-001 breaching', '3.50%', '$133,610'),
    ('5', '$7,214,889', '9.50%', 'LN-002 breaching', '9.75%', '$703,452'),
]
for row in pool_rows:
    cells = table2.add_row().cells
    for i, v in enumerate(row):
        cells[i].text = v
d.add_paragraph("Total pooled reserve across all five pools: $1,320,509.")

d.add_heading('4. Individually Evaluated Loan Analysis', level=1)
d.add_paragraph(
    "LN-008 (Palmetto Print & Packaging Inc.) carries an Internal Risk Rating of 6 and is 97 days past "
    "due, meeting both individual-evaluation triggers under Policy Section 4. The loan is collateral-"
    "dependent: the borrower filed Chapter 11 on December 9, 2025, and repayment is expected to be provided "
    "through liquidation of the manufacturing equipment pledged as collateral."
)
table3 = d.add_table(rows=1, cols=2)
table3.style = 'Light Grid Accent 1'
hdr3 = table3.rows[0].cells
hdr3[0].text, hdr3[1].text = 'Line Item', 'Amount'
li_rows = [
    ('Outstanding Balance', '$1,087,640'),
    ('Appraised Collateral Value (Dec. 2025 inspection)', '$830,000'),
    ('Estimated Disposition Cost (10%)', '$83,000'),
    ('Net Realizable Collateral Value', '$747,000'),
    ('Specific Reserve (Balance − Net Realizable Value)', '$340,640'),
]
for row in li_rows:
    cells = table3.add_row().cells
    cells[0].text, cells[1].text = row
d.add_paragraph(
    "This specific reserve requirement is corroborated by Exception EXC-010 in the Exception Tracking Log, "
    "which independently records a loss reserve as active for this loan, and by the Appraisal and "
    "Collateral Summary note referencing the Chapter 11 filing. No other loan in the portfolio is rated 6 "
    "or 90+ days past due; no other loan requires individual evaluation."
)

d.add_heading('5. Covenant Breach and Watch List Cross-Reference', level=1)
d.add_paragraph(
    "Five loans carry an active covenant breach as of the review date: LN-001, LN-002, LN-004, LN-006, and "
    "LN-010 (LN-008 is excluded from this pool-level analysis because it is individually evaluated in "
    "Section 4, not pooled). LN-001, LN-002, and LN-008 additionally carry a Watch List designation under "
    "the Loan Portfolio Review Policy (Internal Risk Rating of 4 or higher). The Watch List designation "
    "itself does not directly change the pooled reserve calculation — the Risk Rating pool already reflects "
    "elevated risk through its base historical loss rate — but the presence of a breach within a Watch List "
    "loan's pool is what triggers the Q-Factor add-on documented in Section 3. LN-004, LN-006, and LN-010 "
    "are not on the Watch List (their Internal Risk Ratings are 3, below the Rating ≥ 4 threshold), but "
    "their breaches still trigger the Pool 3 Q-Factor because the Q-Factor is a pool-level, not loan-level, "
    "adjustment."
)

d.add_heading('6. Reserve Adequacy Conclusion', level=1)
table4 = d.add_table(rows=1, cols=2)
table4.style = 'Light Grid Accent 1'
hdr4 = table4.rows[0].cells
hdr4[0].text, hdr4[1].text = 'Line Item', 'Amount'
concl_rows = [
    ('Total Pooled Reserve (Ratings 1–5)', '$1,320,509'),
    ('Total Specific Reserve (LN-008)', '$340,640'),
    ('Total Required Reserve', '$1,661,149'),
    ('Currently Booked ALLL Balance', '$1,250,000'),
    ('Reserve Shortfall', '$411,149'),
]
for row in concl_rows:
    cells = table4.add_row().cells
    cells[0].text, cells[1].text = row
d.add_paragraph(
    "The $411,149 shortfall represents 32.9% of the currently booked balance, well outside the ±10% "
    "adequacy tolerance band in Policy Section 6. The reserve is therefore inadequate, and an additional "
    "provision of $411,149 should be recognized in the current quarter and reported to the Board Audit "
    "Committee. The shortfall is driven primarily by the Rating 5 pool (LN-002, $703,452 of required "
    "reserve on a single loan) and the LN-008 specific reserve ($340,640), both of which reflect "
    "deterioration — a payment demand and dual covenant breach on LN-002, and a Chapter 11 filing with "
    "declining collateral value on LN-008 — that the currently booked balance does not appear to fully "
    "capture."
)

d.add_heading('7. Policy/Source Reconciliation', level=1)
d.add_paragraph(
    "Historical loss rates, the disposition cost assumption, the Q-Factor rate, and the currently booked "
    "ALLL balance were taken directly from ALLL_Loss_Rate_and_Reserve_Data.xlsx as pre-calculated policy "
    "inputs; no source conflicts were found among these figures. Risk Rating pool composition was "
    "independently re-tallied from Loan_Register_Q4_2025.xlsx rather than assumed, and the six segments "
    "were confirmed to sum to the full 29-loan, $134,656,509 portfolio."
)
d.add_paragraph(
    "Covenant breach status feeding the Q-Factor determination was independently recomputed rather than "
    "read from the Covenant Compliance Workbook's pre-calculated actual-vs-threshold fields: DSCR was "
    "recomputed from Borrower_Financial_Statements_Summary.xlsx (Net Operating Income ÷ Total Annual Debt "
    "Service), LTV from the Loan Register outstanding balance divided by the Appraisal and Collateral "
    "Summary appraised value, and Current Ratio from Current Assets ÷ Current Liabilities in the financial "
    "statements file. All recomputed values matched the workbook's pre-calculated figures within normal "
    "rounding tolerance (±0.01x or ±0.1 percentage point), and no conflicts were identified between the "
    "pre-calculated workbook values and the independently derived calculations."
)
d.add_paragraph(
    "The LN-008 specific reserve calculation used the most recent appraised collateral value on file "
    "(the December 2025 collateral inspection in the Appraisal and Collateral Summary, $830,000) rather "
    "than an older or origination-date value, consistent with Policy Section 4's requirement to use the "
    "most recent appraisal regardless of loan origination date. This treatment is corroborated by Exception "
    "EXC-010, which independently documents an active loss reserve for this loan in the Exception Tracking "
    "Log."
)

path = os.path.join(OUTDIR, 'Tri_County_ALLL_Reserve_Adequacy_Memo_Q4_2025.docx')
d.save(path)
print('Saved', path)
