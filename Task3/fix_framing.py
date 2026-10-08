from docx import Document

d = Document('check_golden/Loan_Portfolio_Review_Report_Q4_2025.docx')

def replace_paragraph_text(p, new_text):
    for r in p.runs:
        r.text = ''
    if p.runs:
        p.runs[0].text = new_text
    else:
        p.add_run(new_text)

p4 = d.paragraphs[4]
replace_paragraph_text(
    p4,
    "This report documents the Q4 2025 quarterly review of Tri-County Community Bank's commercial loan "
    "portfolio, performed by the Credit Analysis Department to verify compliance with the bank's own "
    "board-approved Loan Portfolio Review Policy. The portfolio under review comprises 29 active commercial "
    "loans with an aggregate outstanding balance of $134,656,509 as of the review date."
)

p54 = d.paragraphs[54]
replace_paragraph_text(
    p54,
    "This section documents how all reported findings were derived and cross-checked against the underlying "
    "source files in accordance with the Loan Portfolio Review Policy requirement to verify pre-calculated "
    "fields against source data and resolve any source conflicts by applying the policy strictly. This "
    "reconciliation supports the report's conclusions on the portfolio's compliance with policy and will be "
    "retained in the credit file to support the next quarterly review."
)

d.save('check_golden/Loan_Portfolio_Review_Report_Q4_2025.docx')
print('saved')
