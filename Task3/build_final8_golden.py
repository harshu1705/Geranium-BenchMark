import docx
import copy
import shutil, os

SRC = os.path.join(os.path.dirname(__file__), "FINAL7_golden", "Tri_County_ALLL_Reserve_Adequacy_Memo_Q4_2025.docx")
OUT_DIR = os.path.join(os.path.dirname(__file__), "FINAL8_golden")
os.makedirs(OUT_DIR, exist_ok=True)
OUT = os.path.join(OUT_DIR, "Tri_County_ALLL_Reserve_Adequacy_Memo_Q4_2025.docx")

d = docx.Document(SRC)

def set_text(idx, new_text):
    p = d.paragraphs[idx]
    assert len(p.runs) == 1, f"paragraph {idx} has {len(p.runs)} runs"
    p.runs[0].text = new_text

def insert_paragraph_after(paragraph, text, style=None):
    new_p_elem = copy.deepcopy(paragraph._p)
    paragraph._p.addnext(new_p_elem)
    new_para = docx.text.paragraph.Paragraph(new_p_elem, paragraph._parent)
    # clear all runs from the copy, then add a single run with new text
    for r in list(new_para.runs):
        r._r.getparent().remove(r._r)
    new_para.add_run(text)
    if style:
        new_para.style = style
    return new_para

# --- Executive Summary ---
set_text(4,
    "Auditors are going to want to see this before quarter close, so here is where the reserve "
    "actually stands. The commercial book runs 29 loans, $134,656,509 outstanding in total, and "
    "under the Bank's ALLL Policy the required reserve against that book comes to $1,871,603: "
    "$1,530,963 pooled across Risk Rating tiers 1 through 5, plus a $340,640 specific reserve "
    "carried against LN-008 (Palmetto Print & Packaging Inc.), which is individually evaluated "
    "because it is collateral-dependent following its Chapter 11 filing. All balances, ratings, "
    "and covenant figures below reflect the portfolio's status as of the December 31, 2025 "
    "quarter-end close, the basis for this quarter's ALLL test."
)

set_text(5,
    "Booked ALLL sits at $1,250,000. That leaves a shortfall of $621,603, or about 33.2% of the "
    "$1,871,603 total required reserve, well outside the ±10%-of-required-reserve band Policy "
    "Section 6 allows. An additional provision of $621,603 should go in this quarter, and the "
    "Board Audit Committee needs to see it."
)

# --- Pooled Reserve Calculation narrative ---
set_text(12,
    "Three loans in Pool 3 are breaching with nothing on file to offset them: LN-004 (Coastal Auto "
    "Group, DSCR 1.14x against a 1.20x floor), LN-006 (Harbor Grille Restaurant Group, Current Ratio "
    "1.03x against 1.20x), and LN-010 (U-Save Self Storage LLC, LTV 85.2% against a 75% cap). None "
    "has a waiver, so the Q-Factor applies to the whole pool. Pool 5 has one breaching loan in the "
    "same boat: LN-002 (Corrigan Retail Partners), LTV 88.0% against 75% and Current Ratio 0.94x "
    "against 1.00x, no waiver there either, so Pool 5 gets the add-on too. Pool 2 turned up a breach "
    "as well, and this one wasn't already flagged anywhere: LN-019 (Interstate Automotive LLC) "
    "recomputes to an 88.0% LTV ($4,221,780 outstanding against the $4,800,000 December 2025 "
    "floor-plan revaluation in Appraisal_and_Collateral_Summary_Q4_2025.xlsx), over its 80% cap. "
    "The Covenant Compliance Workbook still shows LN-019's Actual LTV at a stale 72.8%, calculated "
    "off the prior $5,800,000 figure, and there is no LN-019 entry anywhere in "
    "Exception_Tracking_Log_Q4_2025.xlsx. Nothing waives it, so the full $87,998,690 Pool 2 balance "
    "gets the 25bps add-on too."
)

set_text(13,
    "Pool 4 is different. Its only loan, LN-001 (Talbot Steel Fabricators LLC), is breaching DSCR at "
    "1.07x versus the 1.20x minimum, but Exception EXC-001 shows a waiver approved November 25, 2025 "
    "and running through March 31, 2026, active and not expired as of this review. Policy Section 5 "
    "is explicit that a breach \"offset by a documented, approved waiver\" doesn't trigger the "
    "Q-Factor, so Pool 4 stays at its base rate. Pool 1 isn't breaching anything and doesn't get an "
    "add-on."
)

set_text(14, "Pooled reserve across the five tiers comes to $1,530,963.")

# Table 1 (Pooled Reserve Calculation) - update Risk Rating 2 row
t1 = d.tables[1]
row2 = t1.rows[2]  # header, rating1, rating2 -> index 2
assert row2.cells[0].text == "2"
row2.cells[3].text = "LN-019 breaching, no waiver"
row2.cells[4].text = "0.60%"
row2.cells[5].text = "$527,992"

# --- Covenant Breach and Watch List Cross-Reference ---
set_text(21,
    "Six loans are carrying an active breach right now: LN-001, LN-002, LN-004, LN-006, LN-010, and "
    "LN-019. LN-008 isn't part of this list because it's individually evaluated in Section 4, not "
    "pooled. Each one's Exception Tracking Log entry, where one exists, was pulled and checked for a "
    "Waiver Issued flag and expiry date before any Q-Factor got applied. Only LN-001 has an approved "
    "waiver still in force; LN-002, LN-004, LN-006, and LN-010 show \"Waiver Issued: No\" on their "
    "exceptions, and LN-019 has no exception record at all. All five without an active waiver stay "
    "subject to the add-on."
)

set_text(22,
    "Of those six, only LN-002 carries an explicit \"Watch List\" designation: both of its exceptions, "
    "EXC-002 and EXC-003, record its status as \"Open – Escalated to Watch List.\" No other "
    "exception in Exception_Tracking_Log_Q4_2025.xlsx uses that language (LN-001's is tracked as "
    "\"Open – Monitoring,\" and LN-008's as \"Open – Legal Hold\"), and none of the provided "
    "source files define Watch List status as a function of Internal Risk Rating, so that designation "
    "isn't extended to LN-001 or LN-008 here. Watch List status, where it appears, is an "
    "exception-tracking designation recorded in the Exception Tracking Log; none of the source files "
    "say who assigns it. It's a separate flag from the pool-level Q-Factor and doesn't itself change "
    "the pooled math. LN-004, LN-006, and LN-010 carry no Watch List designation in the exception log "
    "at all, but their breaches still pull the Pool 3 Q-Factor because that add-on applies to the "
    "whole pool, not just the specific loan that's breaching. LN-019 is a further step removed: it "
    "has no entry in Exception_Tracking_Log_Q4_2025.xlsx at all, Watch List or otherwise. Its breach "
    "only surfaced because DSCR, LTV, and Current Ratio were rebuilt from the underlying loan, "
    "financial, and appraisal files rather than pulled from the Covenant Compliance Workbook or the "
    "exception log, and it still pulls the full Pool 2 Q-Factor for the same reason LN-004, LN-006, "
    "and LN-010 do."
)

# --- Reserve Adequacy Conclusion ---
set_text(24,
    "$621,603 works out to about 33.2% of the $1,871,603 total required reserve, well past the "
    "±10%-of-required-reserve band Policy Section 6 allows. That makes the reserve inadequate as "
    "booked, and an additional provision of $621,603 needs to go in this quarter and get reported up "
    "to the Board Audit Committee. Most of the gap traces back to three things: LN-002, where a "
    "single loan in the Rating 5 pool is driving $703,452 of required reserve on its own; LN-008's "
    "$340,640 specific reserve; and the Pool 2 Q-Factor triggered by LN-019's newly recomputed LTV "
    "breach, which alone adds roughly $220,000 across the pool's $87,998,690 balance. All three "
    "reflect deterioration the booked balance hasn't caught up with: a payment demand and two "
    "concurrent covenant breaches on LN-002, a bankruptcy filing and falling collateral value on "
    "LN-008, and a collateral revaluation on LN-019 that neither the Covenant Compliance Workbook nor "
    "the exception log had picked up."
)

# --- Policy/Source Reconciliation ---
set_text(27,
    "The recomputation was not a formality. It caught two live discrepancies. "
    "Covenant_Compliance_Workbook_Q4_2025.xlsx's own Actual LTV and Actual Current Ratio columns for "
    "LN-002 show 71.5% and 1.06x, both comfortably within covenant (75% max LTV, 1.00x minimum "
    "Current Ratio). Recomputing both from the underlying source files tells a different story: LTV "
    "is Outstanding Balance ($7,214,889, per Loan_Register_Q4_2025.xlsx) over Appraised Value "
    "($8,200,000, per the August 14, 2025 appraisal in Appraisal_and_Collateral_Summary_Q4_2025.xlsx), "
    "which is 88.0%, not 71.5%, a breach of the 75% cap. Current Ratio is Current Assets ($813,900) "
    "over Current Liabilities ($865,800), per Borrower_Financial_Statements_Summary.xlsx, which is "
    "0.94x, not 1.06x, a breach of the 1.00x floor. The recomputed Current Ratio lines up exactly "
    "with Exception Tracking Log entry EXC-003, which independently records 0.94x, and the recomputed "
    "LTV is consistent with EXC-002's own note of LTV at 87.9% (the exception log's figure rounds off "
    "a slightly different intermediate calculation than the loan register balance used here; either "
    "way it clears the 75% cap by more than ten points). The workbook's pre-filled Actual columns for "
    "LN-002 are stale, and the recomputed values are the ones that should drive the Q-Factor "
    "determination. Had the workbook's face-value figures been accepted instead, LN-002 would have "
    "appeared covenant-compliant, Pool 5 would have been assessed no Q-Factor, and the total required "
    "reserve would have been understated by roughly $18,000."
)

p27 = d.paragraphs[27]
insert_paragraph_after(p27,
    "The second discrepancy is worse, because nothing in the file set flags it at all. LN-019's "
    "Covenant Compliance Workbook Actual LTV reads 72.8%, which matches an older $5,800,000 "
    "collateral figure and isn't, on its face, a breach. But Appraisal_and_Collateral_Summary_Q4_2025.xlsx "
    "and Loan_Register_Q4_2025.xlsx both carry a December 5, 2025 floor-plan revaluation of $4,800,000 "
    "for the same collateral, and $4,221,780 outstanding against that figure is an 88.0% LTV, over the "
    "80% cap. LN-019 has no entry anywhere in Exception_Tracking_Log_Q4_2025.xlsx, so there was no "
    "exception record to cross-check against and no shortcut around doing the recomputation. Skipping "
    "it would have left Pool 2 at its base 0.35% rate and understated the total required reserve by a "
    "further $219,997. Between LN-002 and LN-019, accepting the pre-calculated fields at face value "
    "would have understated the required reserve by roughly $238,000 combined. DSCR came from Net "
    "Operating Income over Total Annual Debt Service in Borrower_Financial_Statements_Summary.xlsx; "
    "LTV came from the Loan Register balance divided by the appraised value in the Appraisal and "
    "Collateral Summary; Current Ratio came from Current Assets over Current Liabilities in the "
    "financial statements file. Every other loan's recomputed DSCR, LTV, and Current Ratio landed "
    "within normal rounding of the workbook's own figures. LN-002 and LN-019 were the only two "
    "discrepancies found."
)

# Re-fetch paragraphs since document structure changed; find the waiver paragraph by content
for p in d.paragraphs:
    if p.text.startswith("Breach status by itself wasn't treated as enough"):
        assert len(p.runs) == 1
        p.runs[0].text = (
            "Breach status by itself wasn't treated as enough to justify a Q-Factor, though. Every "
            "breaching loan's Waiver Issued, Waiver Approved Date, and Waiver Expiry Date fields in "
            "Exception_Tracking_Log_Q4_2025.xlsx got checked first, consistent with Policy Section 5's "
            "carve-out for breaches \"offset by a documented, approved waiver.\" That check is what "
            "surfaced LN-001's waiver: EXC-001, Waiver Issued marked Yes, approved November 25, 2025 "
            "and running to March 31, 2026, already in force well before the December 31, 2025 review "
            "date, so Pool 4 was reserved at base rate with no add-on. LN-002, LN-004, LN-006, and "
            "LN-010 came back with \"Waiver Issued: No\" on their respective exceptions and were "
            "assessed the Q-Factor accordingly. LN-019 has no exception on file at all, waived or "
            "otherwise, so there was nothing to offset its breach either."
        )
        break
else:
    raise RuntimeError("waiver paragraph not found")

d.save(OUT)
print("saved", OUT)
