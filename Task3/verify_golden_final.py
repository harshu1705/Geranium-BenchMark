from docx import Document

d = Document('check_golden/Tri_County_ALLL_Reserve_Adequacy_Memo_Q4_2025.docx')

# 1. Pool table sums to portfolio total
t0 = d.tables[0]
total_bal = 0
for row in t0.rows[1:]:
    total_bal += int(row.cells[2].text.replace('$', '').replace(',', ''))
print('Pool table balances sum:', total_bal, 'expect 134656509', total_bal == 134656509)

count_sum = sum(int(row.cells[1].text) for row in t0.rows[1:])
print('Pool table counts sum:', count_sum, 'expect 29', count_sum == 29)

# 2. Reserve table math
t1 = d.tables[1]
reserve_sum = 0
for row in t1.rows[1:]:
    reserve_sum += int(row.cells[5].text.replace('$', '').replace(',', ''))
print('Pooled reserve sum:', reserve_sum, 'expect 1310966', reserve_sum == 1310966)

# 3. Specific reserve table
t2 = d.tables[2]
vals = {row.cells[0].text: row.cells[1].text for row in t2.rows[1:]}
print(vals)
bal = int(vals['Outstanding Balance'].replace('$','').replace(',',''))
appr = int(vals['Appraised Collateral Value (Dec. 2025 inspection)'].replace('$','').replace(',',''))
disp = int(vals['Estimated Disposition Cost (10%)'].replace('$','').replace(',',''))
net = int(vals['Net Realizable Collateral Value'].replace('$','').replace(',',''))
spec = int(vals['Specific Reserve (Balance minus Net Realizable Value)'].replace('$','').replace(',',''))
print('disp check', disp == round(appr*0.10), 'net check', net == appr-disp, 'spec check', spec == bal-net)

total_required = reserve_sum + spec
print('total required reserve:', total_required, 'expect 1651606', total_required == 1651606)
shortfall = total_required - 1250000
print('shortfall:', shortfall, 'expect 401606', shortfall == 401606)
print('shortfall pct of booked:', round(100*shortfall/1250000,1))

# name check
for name in ['ABC Industrial','XYZ Retail','Meridian Office Park','Keystone Storage','Lonestar Data','Cascade Tech']:
    found = any(name in p.text for p in d.paragraphs) or any(name in c.text for t in d.tables for row in t.rows for c in row.cells)
    print('old name present:', name, found)
