from docx import Document

d = Document('check_golden/Tri_County_Loan_Portfolio_Exam_Workpaper_Q4_2025.docx')

p0 = d.paragraphs[0]
p1 = d.paragraphs[1]

def set_text(p, new_text):
    for r in p.runs:
        r.text = ''
    if p.runs:
        p.runs[0].text = new_text
    else:
        p.add_run(new_text)

set_text(p0, 'Tri-County Community Bank — Q4 2025 Loan Portfolio Review Report')
set_text(p1, 'Prepared by: Credit Analysis Department, Tri-County Community Bank')

t0 = d.tables[0]
for row in t0.rows:
    if row.cells[0].text.strip() == 'Term Loan':
        for r in row.cells[2].paragraphs[0].runs:
            r.text = ''
        if row.cells[2].paragraphs[0].runs:
            row.cells[2].paragraphs[0].runs[0].text = '$20,281,745'
        else:
            row.cells[2].text = '$20,281,745'

d.save('check_golden/Loan_Portfolio_Review_Report_Q4_2025.docx')
print('done')
