from docx import Document
d = Document('check_golden/Loan_Portfolio_Review_Report_Q4_2025.docx')
print('P0:', d.paragraphs[0].text)
print('P1:', d.paragraphs[1].text)
t0 = d.tables[0]
total = 0
for row in t0.rows[1:]:
    val = row.cells[2].text.replace('$','').replace(',','')
    total += int(val)
    print(row.cells[0].text, row.cells[2].text)
print('table0 total =', total, 'target 134656509, diff', 134656509-total)
