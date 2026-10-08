# -*- coding: utf-8 -*-
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import zipfile
import datetime
from docx import Document
import openpyxl

BASE = r'C:\Users\user\OneDrive\Desktop\geranium project\Task3'
INPUTS = BASE + r'\FIXED_inputs'
GOLDEN = BASE + r'\FIXED_golden'

print("=" * 70)
print("1. METADATA CHECK")
print("=" * 70)
import os
for f in os.listdir(INPUTS):
    path = os.path.join(INPUTS, f)
    if f.endswith('.docx'):
        d = Document(path)
        bad = 'python-docx' in (d.core_properties.author or '')
        print(f"  {'FAIL' if bad else 'OK  '} {f} author={d.core_properties.author!r}")
    elif f.endswith('.xlsx'):
        with zipfile.ZipFile(path) as z:
            app = z.read('docProps/app.xml').decode('utf-8')
        bad = 'openpyxl' in app.lower()
        print(f"  {'FAIL' if bad else 'OK  '} {f} openpyxl marker: {bad}")

gpath = os.path.join(GOLDEN, 'Tri_County_ALLL_Reserve_Adequacy_Memo_Q4_2025.docx')
gd = Document(gpath)
print(f"  {'FAIL' if 'python-docx' in (gd.core_properties.author or '') else 'OK  '} GOLDEN author={gd.core_properties.author!r}")

print()
print("=" * 70)
print("2. FORMATTING CHECK (headers, freeze panes, formulas, real dates)")
print("=" * 70)

wb = openpyxl.load_workbook(os.path.join(INPUTS, 'ALLL_Loss_Rate_and_Reserve_Data.xlsx'))
ws = wb['Loss Rates']
print("  Loss Rates B1:", ws['B1'].value)
print("  Reserve Assumptions B1:", wb['Reserve Assumptions']['B1'].value)
ws3 = wb['Historical Loss Rate Trend (8Q)']
print("  J2 formula:", ws3['J2'].value)
print("  freeze panes set on all sheets:", all(wb[s].freeze_panes for s in wb.sheetnames))

for fname, datecols in [
    ('Loan_Register_Q4_2025.xlsx', ['F2', 'G2', 'L2']),
    ('Borrower_Financial_Statements_Summary.xlsx', ['C2', 'H2']),
    ('Appraisal_and_Collateral_Summary_Q4_2025.xlsx', ['D2']),
    ('Exception_Tracking_Log_Q4_2025.xlsx', ['E2', 'J2']),
]:
    wbx = openpyxl.load_workbook(os.path.join(INPUTS, fname))
    wsx = wbx.active
    print(f"  {fname}: freeze={wsx.freeze_panes}", end=' ')
    for c in datecols:
        v = wsx[c].value
        print(f"{c}={v!r}({type(v).__name__})", end=' ')
    print()

print()
print("=" * 70)
print("3. COVENANT WORKBOOK FLIP CHECK")
print("=" * 70)
wbc = openpyxl.load_workbook(os.path.join(INPUTS, 'Covenant_Compliance_Workbook_Q4_2025.xlsx'), data_only=True)
wsc = wbc.active
row = [wsc.cell(row=3, column=c).value for c in range(1, 9)]
print("  LN-002 row (workbook face value):", row)

# True recompute from source files (unchanged)
wbb = openpyxl.load_workbook(os.path.join(INPUTS, 'Borrower_Financial_Statements_Summary.xlsx'), data_only=True)
wsb = wbb.active
wba = openpyxl.load_workbook(os.path.join(INPUTS, 'Appraisal_and_Collateral_Summary_Q4_2025.xlsx'), data_only=True)
wsa = wba.active
wbl = openpyxl.load_workbook(os.path.join(INPUTS, 'Loan_Register_Q4_2025.xlsx'), data_only=True)
wsl = wbl.active

ln002_fin = [r for r in wsb.iter_rows(min_row=2, values_only=True) if r[0] == 'LN-002'][0]
noi, ds, ca, cl = ln002_fin[3], ln002_fin[4], ln002_fin[5], ln002_fin[6]
ln002_appr = [r for r in wsa.iter_rows(min_row=2, values_only=True) if r[0] == 'LN-002'][0]
appraised = ln002_appr[4]
ln002_reg = [r for r in wsl.iter_rows(min_row=2, values_only=True) if r[0] == 'LN-002'][0]
balance = ln002_reg[4]

true_dscr = noi / ds
true_ltv = balance / appraised * 100
true_cr = ca / cl
print(f"  TRUE recomputed: DSCR={true_dscr:.2f}  LTV={true_ltv:.1f}%  CurrentRatio={true_cr:.2f}")
print(f"  Workbook shows:  DSCR={row[3]}  LTV={row[5]}%  CurrentRatio={row[7]}")
print(f"  Workbook LTV/CR are FALSIFIED (compliant-looking) vs TRUE breaching values: {row[5] != round(true_ltv,1) and row[7] != round(true_cr,2)}")

print()
print("=" * 70)
print("4. GOLDEN MEMO CONSISTENCY CHECK")
print("=" * 70)
full = '\n'.join(p.text for p in gd.paragraphs)
for t in gd.tables:
    for r in t.rows:
        for c in r.cells:
            full += '\n' + c.text

checks = {
    '$134,656,509 (portfolio total)': '$134,656,509' in full,
    '$1,651,606 (total required reserve)': '$1,651,606' in full,
    '$401,606 (shortfall)': '$401,606' in full,
    '24.3% (corrected shortfall %)': '24.3%' in full,
    'No stale 32%': '32%' not in full,
    'No fabricated Watch List rating-4 rule': 'risk rating 4 or higher' not in full and 'Internal Risk Rating of 4 or higher' not in full,
    'No Policy Section 8 / Approval&Distribution': 'Policy Section 8' not in full and 'Approval and Distribution' not in full,
    'LN-002 discrepancy documented': 'stale' in full and '71.5%' in full,
    '$703,452 (Pool 5 reserve, unchanged)': '$703,452' in full,
    '$340,640 (LN-008 specific reserve, unchanged)': '$340,640' in full,
    'Exactly 7 numbered sections': sum(1 for p in gd.paragraphs if p.style.name == 'Heading 1' and p.text[0].isdigit()) == 7,
}
for k, v in checks.items():
    print(f"  {'PASS' if v else 'FAIL'}  {k}")
