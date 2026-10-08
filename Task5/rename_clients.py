# -*- coding: utf-8 -*-
import docx, os, shutil

REPLACEMENTS = [
    ('Whitfield', 'Pruitt'),
    ('Gerald', 'Daniel'),
    ('Linda', 'Renee'),
]

def replace_in_doc(path):
    d = docx.Document(path)
    changed = False
    for p in d.paragraphs:
        for run in p.runs:
            new_text = run.text
            for old, new in REPLACEMENTS:
                new_text = new_text.replace(old, new)
            if new_text != run.text:
                run.text = new_text
                changed = True
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        new_text = run.text
                        for old, new in REPLACEMENTS:
                            new_text = new_text.replace(old, new)
                        if new_text != run.text:
                            run.text = new_text
                            changed = True
    if changed:
        d.save(path)
    return changed

BASE = os.path.join(os.path.dirname(__file__), 'PIVOT3_inputs')
for f in os.listdir(BASE):
    if f.endswith('.docx'):
        c = replace_in_doc(os.path.join(BASE, f))
        print(f, 'changed:', c)

GOLD_DIR = os.path.join(os.path.dirname(__file__), 'PIVOT3_golden')
old_golden = os.path.join(GOLD_DIR, 'Whitfield_Retirement_Readiness_Assessment.docx')
new_golden = os.path.join(GOLD_DIR, 'Pruitt_Retirement_Readiness_Assessment.docx')
c = replace_in_doc(old_golden)
print('golden changed:', c)
shutil.move(old_golden, new_golden)
print('renamed golden to', new_golden)
