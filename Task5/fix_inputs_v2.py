# -*- coding: utf-8 -*-
"""Reviewer-driven corrections to the Pruitt input files.

1. Social_Security_Benefit_Estimates.docx
   - age-65 amounts now follow the correct ~13.3% early-claiming reduction for FRA 67
     (was ~9%, which is wrong for anyone born 1960 or later)
   - the FRA section states the actual reduction formula
   - the spousal paragraph no longer falsely dismisses a possible top-up for Renee
2. Client_Intake_and_Goals_Summary.docx
   - documents how the firm applies its 4% screen when a self-funded bridge
     period precedes the start of Social Security
3. Household_Budget_and_Planning_Assumptions.xlsx
   - housing note no longer contradicts the intake's "owns the home outright"
   - health-insurance bridge line states it is already in age-62..64 dollars
4. Insurance_and_Estate_Summary.docx
   - aligns the $1,800 quote with the budget workbook's treatment
"""
import os
import copy
import docx
import openpyxl

BASE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "PIVOT3_inputs")


def set_text(par, text):
    """Replace a paragraph's text, keeping the first run's formatting."""
    if par.runs:
        par.runs[0].text = text
        for r in par.runs[1:]:
            r.text = ""
    else:
        par.add_run(text)


def clone_after(par, text):
    """Insert a new paragraph with the same style directly after `par`."""
    new = copy.deepcopy(par._p)
    par._p.addnext(new)
    newpar = docx.text.paragraph.Paragraph(new, par._parent)
    set_text(newpar, text)
    return newpar


# ============================================================
# 1. Social Security Benefit Estimates
# ============================================================
path = os.path.join(BASE, "Social_Security_Benefit_Estimates.docx")
d = docx.Document(path)
paras = d.paragraphs

REPL = {
    "Age 65: $3,100 per month":
        "Age 65: $2,947 per month",
    "Age 65: $1,450 per month":
        "Age 65: $1,387 per month",
}

HOW = (
    "Each estimate is based on the worker's Average Indexed Monthly Earnings (AIME), calculated from the "
    "highest 35 years of wage-indexed earnings on record, which is then run through the Social Security "
    "Primary Insurance Amount (PIA) formula to arrive at the benefit payable at Full Retirement Age. The "
    "age-62 and age-65 figures above are that same PIA reduced for early claiming as described below, and "
    "the age-67 figure is the unreduced PIA itself. Dollar amounts are rounded planning estimates drawn "
    "from the earnings record and may differ by a few dollars from a literal percentage of the age-67 "
    "figure. These are estimates only; actual benefits depend on continued earnings between now and the "
    "claiming date and are not guaranteed."
)

FRA = (
    "For anyone born in 1960 or later, including both Daniel and Renee, Full Retirement Age (FRA) under "
    "current law is 67. A benefit claimed before FRA is permanently reduced by five-ninths of one percent "
    "for each of the first 36 months it is taken early, plus five-twelfths of one percent for each "
    "additional month beyond 36. Claiming at 62, the earliest eligibility age, is 60 months early and "
    "produces a permanent reduction of approximately 30% from the FRA benefit. Claiming at 65 is 24 months "
    "early and produces a permanent reduction of approximately 13.3% from the FRA benefit. Note that the "
    "age-65 reduction is larger than a straight-line interpolation between the age-62 and age-67 figures "
    "would suggest, because the reduction accrues per month rather than per year of delay and the steeper "
    "five-ninths rate applies across the whole 24-month window. Waiting past FRA accrues delayed "
    "retirement credits of approximately 8% per year up to age 70, after which no further credits accrue; "
    "post-FRA amounts are not shown in this statement. The figures shown above for ages 62, 65, and 67 "
    "already reflect these respective reductions."
)

SPOUSAL = (
    "The estimates above reflect only each individual's own earnings record. Spousal benefits (up to 50% "
    "of a spouse's Full Retirement Age benefit) and survivor benefits (up to 100% of a deceased spouse's "
    "benefit) may be available under separate rules and are not included in the figures above. One half of "
    "Daniel's Full Retirement Age benefit is $1,700 per month, modestly above Renee's own Full Retirement "
    "Age benefit of $1,600, so a spousal top-up of up to roughly $100 per month could be payable to Renee "
    "at her Full Retirement Age. Any such top-up would be payable only once Daniel has filed, and would "
    "itself be reduced if Renee claimed before her own Full Retirement Age, so the amount actually payable "
    "in any given claiming scenario would be smaller than that and could be zero. The figures in this "
    "statement exclude it. Because the potential amount is small relative to this household's projected "
    "income need and cannot be confirmed without a determination from the Social Security Administration, "
    "it should be treated as possible upside rather than planned income, and eligibility should be "
    "confirmed with the Social Security Administration directly at the time of filing."
)

hits = 0
for par in paras:
    t = par.text.strip()
    if t in REPL:
        set_text(par, REPL[t])
        hits += 1
    elif t.startswith("Each estimate is based on the worker's Average Indexed"):
        set_text(par, HOW)
        hits += 1
    elif t.startswith("For anyone born in 1960 or later"):
        set_text(par, FRA)
        hits += 1
    elif t.startswith("The estimates above reflect only each individual's own earnings record."):
        set_text(par, SPOUSAL)
        hits += 1

assert hits == 5, "expected 5 Social Security edits, made {}".format(hits)
d.save(path)
print("[ok] Social_Security_Benefit_Estimates.docx  ({} edits)".format(hits))

# ============================================================
# 2. Client Intake and Goals Summary
# ============================================================
path = os.path.join(BASE, "Client_Intake_and_Goals_Summary.docx")
d = docx.Document(path)

BRIDGE_RULE = (
    "Applying the 4% screen when a self-funded bridge period comes first: where a plan has the client "
    "living entirely off the portfolio for some years before Social Security begins, the screen is applied "
    "at the point ongoing steady-state withdrawals begin, that is, to the projected portfolio balance on "
    "the date Social Security income starts, with each bridge year's drawdown modeled explicitly year by "
    "year beforehand. Treating a temporary, elevated bridge-year withdrawal as though it were the "
    "permanent withdrawal rate would misstate the sustainability of what is by definition a temporary "
    "funding gap. The 30-year horizon continues to be measured from the retirement date, so applying the "
    "4% rate at a later screen date is conservative relative to the horizon actually remaining at that "
    "date."
)

anchor = None
for par in d.paragraphs:
    if par.text.strip().startswith("Sustainable withdrawal planning uses an initial 4% withdrawal rate"):
        anchor = par
        break
assert anchor is not None, "4% assumption bullet not found"
clone_after(anchor, BRIDGE_RULE)
d.save(path)
print("[ok] Client_Intake_and_Goals_Summary.docx  (bridge-period screening rule added)")

# ============================================================
# 3. Household Budget and Planning Assumptions
# ============================================================
path = os.path.join(BASE, "Household_Budget_and_Planning_Assumptions.xlsx")
wb = openpyxl.load_workbook(path)

ws = wb["Current vs Retirement Budget"]
fixed_housing = False
for row in ws.iter_rows(min_row=2):
    if row[0].value and str(row[0].value).startswith("Housing"):
        row[3].value = ("Residence owned outright, no mortgage; property taxes, insurance and "
                        "maintenance continue and are expected to rise")
        fixed_housing = True
assert fixed_housing, "housing row not found"

ws = wb["Bridge and Contribution Detail"]
fixed_bridge = False
for row in ws.iter_rows(min_row=2):
    if row[0].value and str(row[0].value).startswith("Health insurance bridge"):
        row[3].value = ("Ages 62-64 only, until Medicare eligibility at 65. Figure is the planning cost "
                        "for each of those three bridge years as stated; do not apply the general 3% "
                        "expense inflation assumption to this line")
        fixed_bridge = True
assert fixed_bridge, "bridge row not found"
ws.column_dimensions['D'].width = 60

wb.save(path)
print("[ok] Household_Budget_and_Planning_Assumptions.xlsx  (housing note, bridge note)")

# ============================================================
# 4. Insurance and Estate Summary
# ============================================================
path = os.path.join(BASE, "Insurance_and_Estate_Summary.docx")
d = docx.Document(path)
QUOTE = (
    "A current quote for ACA marketplace or COBRA continuation coverage for the household runs "
    "approximately $1,800 per month combined until Medicare eligibility. This is the figure carried into "
    "the household budget workbook as the health insurance bridge cost, stated there as the expected "
    "monthly cost during each of the ages 62 through 64 bridge years."
)
found = False
for par in d.paragraphs:
    if par.text.strip().startswith("A current quote for ACA marketplace"):
        set_text(par, QUOTE)
        found = True
assert found, "health quote bullet not found"
d.save(path)
print("[ok] Insurance_and_Estate_Summary.docx  (bridge quote basis clarified)")

print("\nAll input corrections applied.")
