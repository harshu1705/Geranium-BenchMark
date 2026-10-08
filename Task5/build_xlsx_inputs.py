# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime
import os

BASE = os.path.join(os.path.dirname(__file__), "PIVOT3_inputs")
os.makedirs(BASE, exist_ok=True)

HEADER_FILL = PatternFill(fill_type='solid', fgColor='1F4E78')
HEADER_FONT = Font(bold=True, color='FFFFFF')

def style_header(ws, ncols, row=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.freeze_panes = ws.cell(row=row + 1, column=1).coordinate

def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

# ============================================================
# 1. Account Statements Summary
# ============================================================
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Account Statements'
headers = ['Account', 'Owner', 'Statement Date', 'Balance ($)', 'Equity Allocation (%)', 'Fixed Income Allocation (%)', 'Cash Allocation (%)']
ws.append(headers)
rows = [
    ['Gerald 401(k)', 'Gerald', datetime(2026,8,31), 680000, 0.65, 0.30, 0.05],
    ['Linda 403(b)', 'Linda', datetime(2026,8,31), 310000, 0.70, 0.25, 0.05],
    ['Joint Taxable Brokerage', 'Joint', datetime(2026,8,31), 215000, 0.80, 0.20, 0.00],
    ['Linda Roth IRA', 'Linda', datetime(2026,8,31), 95000, 0.90, 0.10, 0.00],
    ['Joint Savings', 'Joint', datetime(2026,8,31), 45000, 0.00, 0.00, 1.00],
]
for r in rows:
    ws.append(r)
total = sum(r[3] for r in rows)
ws.append(['Total', '', '', total, '', '', ''])
style_header(ws, len(headers))
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=4, max_col=4):
    for cell in row:
        cell.number_format = '#,##0'
for row in ws.iter_rows(min_row=2, max_row=ws.max_row-1, min_col=5, max_col=7):
    for cell in row:
        cell.number_format = '0%'
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=3, max_col=3):
    for cell in row:
        if isinstance(cell.value, datetime):
            cell.number_format = 'mm/dd/yyyy'
autosize(ws, [26, 10, 16, 14, 16, 18, 14])

ws2 = wb.create_sheet('Taxable Brokerage Detail')
ws2.append(['Item', 'Value ($)'])
ws2.append(['Joint Taxable Brokerage - Current Market Value', 215000])
ws2.append(['Joint Taxable Brokerage - Aggregate Cost Basis', 150000])
ws2.append(['Unrealized Gain', 65000])
style_header(ws2, 2)
for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row, min_col=2, max_col=2):
    for cell in row:
        cell.number_format = '#,##0'
autosize(ws2, [42, 16])

wb.save(os.path.join(BASE, 'Account_Statements_Summary.xlsx'))
print('total assets', total)

# ============================================================
# 2. Household Budget and Planning Assumptions
# ============================================================
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Current vs Retirement Budget'
headers = ['Expense Category', 'Current Monthly Spend ($)', 'Expected Retirement Monthly Spend ($)', 'Notes']
ws.append(headers)
rows = [
    ['Housing (mortgage, taxes, insurance, maintenance)', 2400, 2600, 'Mortgage paid off in 3 years; taxes/insurance/maintenance continue and rise'],
    ['Healthcare (premiums, out-of-pocket)', 650, 1200, 'Expected to rise; see health insurance bridge cost separately for pre-Medicare years'],
    ['Food and household', 900, 850, 'Modest reduction expected'],
    ['Transportation', 700, 500, 'One commute eliminated; keeping both vehicles'],
    ['Travel and leisure', 500, 1300, 'Increased travel is a stated retirement priority'],
    ['Insurance (life, auto, umbrella)', 380, 380, 'Term life premium continues until policy expiration at Gerald age 68'],
    ['Retirement account contributions', 2800, 0, 'Contributions stop at retirement'],
    ['Brokerage contributions', 300, 0, 'Contributions stop at retirement'],
    ['Miscellaneous / discretionary', 570, 1670, 'Includes gifts, hobbies, home projects'],
]
for r in rows:
    ws.append(r)
current_total = sum(r[1] for r in rows)
retirement_total = sum(r[2] for r in rows)
ws.append(['Total Monthly', current_total, retirement_total, ''])
ws.append(['Total Annual', current_total*12, retirement_total*12, ''])
style_header(ws, len(headers))
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=2, max_col=3):
    for cell in row:
        cell.number_format = '#,##0'
autosize(ws, [42, 24, 30, 46])
print('current monthly', current_total, 'retirement monthly', retirement_total, 'retirement annual', retirement_total*12)

ws2 = wb.create_sheet('Bridge and Contribution Detail')
ws2.append(['Item', 'Monthly Amount ($)', 'Annual Amount ($)', 'Applies'])
ws2.append(['Health insurance bridge (ACA/COBRA quote, combined)', 1800, 21600, 'Ages 62-64 only, until Medicare eligibility at 65'])
ws2.append(['Combined retirement account + brokerage contributions', 3100, 37200, 'Now through retirement date only (24 months)'])
style_header(ws2, 4)
for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row, min_col=2, max_col=3):
    for cell in row:
        cell.number_format = '#,##0'
autosize(ws2, [50, 20, 20, 42])

wb.save(os.path.join(BASE, 'Household_Budget_and_Planning_Assumptions.xlsx'))
print('DONE')
