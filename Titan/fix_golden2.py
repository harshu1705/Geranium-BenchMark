import sys, io, os, copy
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

BASE = os.path.dirname(os.path.abspath(__file__))
P = os.path.join(BASE, 'golden', 'Project_Titan_Due_Diligence_Report.docx')
d = Document(P)

def tables(doc):
    return [Table(c, doc) for c in doc.element.body.iterchildren()
            if c.tag.split('}')[-1] == 'tbl']

def set_cell(cell, text):
    ps = cell.paragraphs
    first = ps[0]
    for extra in ps[1:]:
        extra._p.getparent().remove(extra._p)
    if not first.runs:
        first.add_run('')
    first.runs[0].text = text
    for r in first.runs[1:]:
        r.text = ''

log = []
ts = tables(d)
t0 = ts[0]
t_risk = [t for t in ts if t.rows[0].cells[0].text.strip() == 'Risk Area'][0]

# --- Table 0 R7: Stakeholder Sign-offs finding was CISO-only --------
assert 'Stakeholder Sign-offs' in t0.rows[7].cells[0].text
set_cell(t0.rows[7].cells[2],
    'Two required sign-offs are unobtained and two are conditional. Chief Legal Officer '
    'Dr. Lisa Pham is Pending/BLOCKED awaiting an updated GDPR risk assessment, with no owner or '
    'deadline recorded. Chief Privacy Officer Kevin Park is Pending, awaiting Legal’s view '
    'on the GDPR audit implications, also with no deadline. CISO Priya Deshpande is Conditional '
    'on resolving VULN-0019 before the closing date — the tracker states no calendar date '
    'for that condition, though the GDPR Privacy Audit Checklist separately records a September '
    '2, 2026 target for AC-7/VULN-0019. Head of QA Anita Flores is Conditional on resolving '
    'VULN-0049 within 30 days post-close, which the September 15, 2026 target satisfies. '
    'Separately, Sarah Chen’s Approved status rests on an assertion the pilot data '
    'contradicts. The gate is NOT MET until the two Pending sign-offs are cleared. See §7.')
log.append('Table 0 R7: Stakeholder Sign-offs finding now covers all four non-clean sign-offs')

# --- Risk matrix R5: name the VP and cite §4.5.b -------------------
r5 = t_risk.rows[5]
assert 'STRESS-001' in r5.cells[1].text, r5.cells[1].text
set_cell(r5.cells[4],
    'P99 = 4.2s against a ≤3.0s threshold. Policy §4.5.b is satisfied: VP Engineering '
    'Raj Malhotra gave formal written risk acceptance in the M&A Sign-off Tracker on August 20, '
    '2026. That acceptance supports a Conditional Go only and is premised on the '
    'under-provisioned test environment, so re-confirm it after re-testing at ≥80% capacity. '
    'Query optimization required post-merger. See §6.')
log.append('Risk matrix R5: named Raj Malhotra and cited Policy §4.5.b')

d.save(P)
print('APPLIED')
for i, m in enumerate(log, 1):
    print(f'  {i}. {m}')
