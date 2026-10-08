import sys, io, os, copy
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

BASE = os.path.dirname(os.path.abspath(__file__))
P = os.path.join(BASE, 'golden', 'Project_Titan_Due_Diligence_Report.docx')

d = Document(P)

# ------------------------------------------------------------------
# helpers
# ------------------------------------------------------------------
def body_items(doc):
    """Return ordered list of ('p', Paragraph) / ('tbl', Table)."""
    out = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.split('}')[-1]
        if tag == 'p':
            out.append(('p', Paragraph(child, doc)))
        elif tag == 'tbl':
            out.append(('tbl', Table(child, doc)))
    return out

def paras(doc):
    return [it[1] for it in body_items(doc) if it[0] == 'p']

def tables(doc):
    return [it[1] for it in body_items(doc) if it[0] == 'tbl']

def find_para(doc, needle):
    for p in paras(doc):
        if needle in p.text:
            return p
    raise LookupError('paragraph not found: ' + needle[:60])

def set_runs(para, lead, rest):
    """Rewrite a paragraph as [bold lead run][normal rest run], preserving
    the formatting already carried by the paragraph's first two runs."""
    rs = para.runs
    assert rs, 'paragraph has no runs'
    if rest is None:
        rs[0].text = lead
        for r in rs[1:]:
            r.text = ''
        return
    if len(rs) == 1:
        # clone run 0 so we keep character formatting for the tail
        new = copy.deepcopy(rs[0]._r)
        rs[0]._r.addnext(new)
        rs = para.runs
    rs[0].text = lead
    rs[1].text = rest
    for r in rs[2:]:
        r.text = ''

def set_cell(cell, text):
    """Replace a table cell's text, keeping the first paragraph/run formatting."""
    ps = cell.paragraphs
    first = ps[0]
    for extra in ps[1:]:
        extra._p.getparent().remove(extra._p)
    if not first.runs:
        first.add_run('')
    first.runs[0].text = text
    for r in first.runs[1:]:
        r.text = ''

def cell_text(cell):
    return cell.text

def clone_para_after(anchor, text, lead=None):
    """Insert a copy of `anchor` (same style/format) right after it."""
    new_p = copy.deepcopy(anchor._p)
    anchor._p.addnext(new_p)
    np = Paragraph(new_p, anchor._parent)
    if lead is None:
        set_runs(np, text, None)
    else:
        set_runs(np, lead, text)
    return np

def clone_row(table, src_index):
    """Append a deep copy of an existing row; returns the new row."""
    src = table.rows[src_index]
    new_tr = copy.deepcopy(src._tr)
    src._tr.getparent().append(new_tr)
    return table.rows[-1]

def move_row_after(table, row, after_index):
    table.rows[after_index]._tr.addnext(row._tr)

def move_row_before(table, row, before_index):
    table.rows[before_index]._tr.addprevious(row._tr)

log = []

# ==================================================================
# 1. EXECUTIVE SUMMARY
# ==================================================================
p6 = find_para(d, 'The acquisition of HealthSync Inc. cannot proceed as of this report date')
p6.runs[0].text = p6.runs[0].text.rstrip() + (
    ' Separately, Policy §5 requires every vulnerability with an effective severity of '
    'Blocker to be resolved prior to the Board Review. VULN-0031 is the only effective Blocker '
    'in the scan, it remains Open, and its documented remediation target of September 6, 2026 '
    'falls three days after the September 3, 2026 Board Review — so §5 cannot be '
    'satisfied on the current schedule either.'
)
log.append('P6 exec-summary: appended Policy §5 / Board Review finding')

# Replace the mis-placed AC-7 bullet with a concise CISO-conditional bullet.
p13 = find_para(d, 'AC-7 → VULN-0019 (PARTIALLY MET): GDPR control AC-7')
# (guard: the §4 body paragraph has the same lead-in, so pick the List Bullet one)
for p in paras(d):
    if p.style.name == 'List Bullet' and 'AC-7' in p.text and 'VULN-0019' in p.text:
        p13 = p
        break
set_runs(p13,
    'The CISO’s sign-off (Priya Deshpande) is conditional on resolving VULN-0019 before the '
    'closing date. The M&A Sign-off Tracker states no calendar date for that condition; the GDPR '
    'Privacy Audit Checklist separately records a September 2, 2026 remediation target for '
    'AC-7/VULN-0019.', None)
log.append('P13 exec-summary: replaced duplicated AC-7 body text with a concise CISO bullet')

# New infrastructure bullet for the STRESS-003 memory breach, placed after the
# existing stress-environment bullet.
p12 = find_para(d, 'The stress-test environment was provisioned below the Policy 4.5.c minimum')
clone_para_after(p12,
    'STRESS-003 (Bulk PII Export) met its throughput SLA at every percentile but peaked at '
    '3.36 GB against a 3.0 GB memory threshold — 112% of the permitted ceiling — with '
    'memory not returning to baseline after the run. This corresponds to VULN-0044.')
log.append('Exec summary: inserted STRESS-003 memory-breach bullet')

# ==================================================================
# 2. POLICY EVALUATION TABLE (Table 0)
# ==================================================================
t0 = tables(d)[0]

# --- §4.5.a : P50/P95 across ALL scenarios -------------------------
assert '4.5.a' in cell_text(t0.rows[3].cells[0]), cell_text(t0.rows[3].cells[0])
set_cell(t0.rows[3].cells[1], 'MET (WITH CAVEATS)')
set_cell(t0.rows[3].cells[2],
    'Every stress scenario that reported numeric percentiles met its declared SLA metric at both '
    'P50 and P95: STRESS-001 0.8s / 1.9s (≤3.0s); STRESS-002 0.9s / 1.6s (≤2.0s); '
    'STRESS-003 680 / 620 rec-min (≥500); STRESS-005 45ms / 120ms (≤200ms); '
    'STRESS-006 8s / 18s (≤30s); STRESS-007 0.3s / 0.7s (≤1.5s); STRESS-008 0.5s / 2.1s '
    '(≤5.0s). STRESS-004 was recorded qualitatively as “stable” rather than as '
    'numeric percentiles, so its P50/P95 compliance cannot be numerically verified from the '
    'source data. Caveats: STRESS-003 cleared its throughput SLA but breached a separate memory '
    'threshold (§6); and every result is provisional under §4.5.c. See §6.')
log.append('Table 0 R3 (§4.5.a): rewrote to assess P50/P95 across all eight scenarios')

# --- §4.5.b : P99 failure + VP Engineering risk acceptance ---------
assert '4.5.b' in cell_text(t0.rows[4].cells[0]), cell_text(t0.rows[4].cells[0])
set_cell(t0.rows[4].cells[1], 'P99 FAIL — §4.5.b CONDITION SATISFIED')
set_cell(t0.rows[4].cells[2],
    'STRESS-001 (Customer Record Search API) is the only P99 breach across the eight scenarios: '
    'P99 = 4.2s against a ≤3.0s threshold. Policy §4.5.b permits a deal with a P99 SLA '
    'failure to proceed only as a CONDITIONAL GO, and only where formal written risk acceptance '
    'is obtained from the VP of Engineering. VP Engineering Raj Malhotra recorded that acceptance '
    'in the M&A Sign-off Tracker on August 20, 2026 (“accepting P99 risk given test env '
    'scale”), so the §4.5.b requirement is satisfied. Note the ceiling this creates: '
    '§4.5.b can support a Conditional Go at best — it cannot cure the §4.4.c NO-GO '
    'trigger. Note also that Mr. Malhotra’s acceptance is expressly premised on the test '
    'environment scale that §4.5.c finds inadequate, so it should be re-confirmed against a '
    'compliant environment.')
log.append('Table 0 R4 (§4.5.b): rewrote with VP-Engineering risk acceptance analysis')

# --- new row: §5 Blocker remediation vs Board Review ---------------
new_row = clone_row(t0, 8)
set_cell(new_row.cells[0], 'Blocker Remediation (§5)')
set_cell(new_row.cells[1], 'NOT MET')
set_cell(new_row.cells[2],
    'Policy §5 requires any vulnerability with an effective severity of Blocker to be '
    'resolved prior to the Board Review. VULN-0031 is the sole effective Blocker among the 47 '
    'scanned vulnerabilities, its status is still Open, and its documented remediation target is '
    'September 6, 2026 — three days after the September 3, 2026 Board Review. §5 '
    'therefore cannot be met on the documented schedule.')
move_row_before(t0, new_row, 8)
log.append('Table 0: inserted §5 Blocker-remediation row (VULN-0031 Sep 6 vs Sep 3 Board Review)')

# ==================================================================
# 3. VULNERABILITY ANALYSIS — §5 assessment paragraph
# ==================================================================
p24 = find_para(d, 'Policy §5 Compliance Assessment')
set_runs(p24, 'Policy §5 Compliance Assessment: ',
    'Policy §5 imposes two distinct duties. First, any vulnerability with an effective '
    'severity of Blocker must be resolved prior to the Board Review. VULN-0031 is the only '
    'effective Blocker; it remains Open and is targeted for September 6, 2026, which is three '
    'days after the September 3, 2026 Board Review — the Blocker duty therefore fails on '
    'the documented schedule. Second, each effective-Critical vulnerability must be either '
    '(a) resolved, or (b) covered by a documented mitigation plan approved by the CTO; the Policy '
    'does not authorize a “formal CTO waiver” as an alternative. There are 5 '
    'effective-Critical vulnerabilities (VULN-0012, VULN-0019, VULN-0023, VULN-0037, VULN-0041). '
    'None is marked Resolved in the Security Vulnerability Scan. The M&A Sign-off Tracker '
    'contains no CTO role and records no CTO-approved mitigation plan for any of them. Policy '
    '§5 is therefore NOT MET on both limbs.')
log.append('§3: expanded Policy §5 assessment to cover the Blocker limb and the Board Review date')

# ==================================================================
# 4. GDPR COMPLIANCE STATUS — SC-10 paragraph
# ==================================================================
p27 = find_para(d, 'SC-10 → VULN-0031 (NOT MET)')
set_runs(p27, 'SC-10 → VULN-0031 (NOT MET): ',
    'GDPR control SC-10 (Automatic Logoff) is rated Not Met. The root cause is VULN-0031 — '
    'the auto session timeout is not enforced on the new executive dashboard. Under Policy '
    '§4.4.c this is an automatic NO-GO for Not Met controls, and no CISO waiver is '
    'documented anywhere in the supplied materials. The VULN-0031 remediation target date is '
    'listed as September 6, 2026. That target falls three days AFTER the September 3, 2026 Board '
    'Review. Policy §5 requires any vulnerability with an effective severity of Blocker to '
    'be resolved prior to the Board Review, and VULN-0031 carries an effective severity of '
    'Blocker (rated Blocker by both the QA Auditor and Target IT). The September 6 target '
    'therefore breaches Policy §5 as well as §4.4.c: on the documented schedule the '
    'Blocker will still be open when the Board convenes. The 30-day post-close remediation window '
    'in §4.4.b is not available here — it applies only to Partially Met controls, not '
    'to a Not Met control.')
log.append('§4 SC-10: added Sep 6 vs Sep 3 Board Review comparison and Policy §5 linkage')

# ==================================================================
# 5. INFRASTRUCTURE / STRESS-TEST ASSESSMENT (§6)
# ==================================================================
p39 = find_para(d, 'STRESS-001 — P99 SLA Failure')
set_runs(p39, 'STRESS-001 — P99 SLA Failure (Policy §4.5.b): ',
    'The Customer Record Search API (STRESS-001) is the only scenario of the eight that failed at '
    'the P99 percentile. The measured P99 response time was 4.2s against an SLA threshold of '
    '≤3.0s — a 1.2s overshoot, or 140% of the permitted ceiling. P50 (0.8s) and P95 '
    '(1.9s) were comfortably within SLA, so the failure is confined to the tail. Policy '
    '§4.5.b governs exactly this situation: where an SLA metric fails at P99, the deal may '
    'proceed as a CONDITIONAL GO only if formal written risk acceptance is obtained from the VP '
    'of Engineering. That acceptance exists. VP Engineering Raj Malhotra signed off in the M&A '
    'Sign-off Tracker on August 20, 2026 with the note “Infrastructure looks acceptable, '
    'accepting P99 risk given test env scale,” which satisfies the §4.5.b requirement '
    'for this scenario. Two qualifications attach. First, §4.5.b is a gate to a Conditional '
    'Go only — it cannot lift the §4.4.c NO-GO trigger, so the overall recommendation '
    'is unaffected. Second, Mr. Malhotra’s acceptance is expressly conditioned on '
    '“test env scale,” i.e. the same under-provisioned environment that fails '
    '§4.5.c below; the acceptance should be re-confirmed once the tests are re-run on a '
    'compliant environment, because the true production P99 is not established by this data. '
    'Query optimization on concurrent search + filter combinations is recommended post-merger.')
log.append('§6 P39: rewrote STRESS-001 P99 analysis with full Policy §4.5.b determination')

# --- new §4.5.a paragraph + supporting table ----------------------
p_a = clone_para_after(p39,
    'Policy §4.5.a requires infrastructure stress tests to meet their target SLA thresholds '
    'at the P50 and P95 percentiles. Assessed against the declared SLA metric for each scenario, '
    'that requirement is met everywhere it can be measured:',
    lead='§4.5.a — P50 / P95 Compliance Across All Scenarios: ')
log.append('§6: inserted §4.5.a P50/P95 assessment paragraph')

# Build the P50/P95 table by cloning the existing pilot table's look.
sla_rows = [
    ('Scenario', 'SLA Metric', 'Threshold', 'P50', 'P95', '§4.5.a Result'),
    ('STRESS-001 Customer Record Search API', 'Response Time', '≤3.0s', '0.8s', '1.9s', 'MET'),
    ('STRESS-002 Executive Dashboard Load', 'Page Load Time', '≤2.0s', '0.9s', '1.6s', 'MET'),
    ('STRESS-003 Bulk PII Export', 'Throughput', '≥500 rec/min', '680 rec/min', '620 rec/min', 'MET'),
    ('STRESS-004 API Gateway — Concurrent Users', 'Concurrent Sessions', '500 users', 'stable', 'stable', 'NOT NUMERICALLY VERIFIABLE'),
    ('STRESS-005 Database Query Latency', 'Avg Query Time', '≤200ms', '45ms', '120ms', 'MET'),
    ('STRESS-006 Report Generation', 'Generation Time', '≤30s', '8s', '18s', 'MET'),
    ('STRESS-007 Login/SSO Flow', 'Auth Response Time', '≤1.5s', '0.3s', '0.7s', 'MET'),
    ('STRESS-008 Real-time Notification Delivery', 'Delivery Latency', '≤5.0s', '0.5s', '2.1s', 'MET'),
]
sla_tbl = d.add_table(rows=len(sla_rows), cols=6)
sla_tbl.style = d.tables[0].style
for ri, row in enumerate(sla_rows):
    for ci, val in enumerate(row):
        cell = sla_tbl.cell(ri, ci)
        para = cell.paragraphs[0]
        run = para.add_run(val)
        run.font.size = Pt(9)
        run.bold = (ri == 0)
# move it directly after the §4.5.a paragraph
p_a._p.addnext(sla_tbl._tbl)
log.append('§6: inserted P50/P95 compliance table for all eight scenarios')

# --- STRESS-003 memory-threshold breach paragraph -----------------
anchor = Paragraph(sla_tbl._tbl.getnext(), d) if sla_tbl._tbl.getnext() is not None else None
p_mem = clone_para_after(p_a,
    'The Bulk PII Export scenario is recorded in the workbook with a status of '
    '“PASS*” — a qualified pass, not a clean one, and the qualification matters. '
    'Its declared SLA metric is throughput, and throughput is comfortably met at every '
    'percentile (680 / 620 / 540 records per minute against a ≥500 records/min threshold), '
    'so §4.5.a and §4.5.b are both satisfied for this scenario. The Resource '
    'Utilization data for the same run, however, records a peak memory figure of 3.36 GB against '
    'a 3.0 GB memory threshold — an overage of 0.36 GB, or 112% of the permitted ceiling '
    '— and notes that the resident set did not return to its pre-run baseline after the run '
    'completed. That signature (sustained growth over a 180-minute window, no return to '
    'baseline) matches VULN-0044, “Memory leak during bulk PII export under sustained '
    'load” (Major, module Bulk PII Export), whose own scan note records a linkage to a '
    'performance-benchmark memory threshold breach occurring only after roughly two hours of '
    'continuous export. STRESS-003 therefore did not pass cleanly: the throughput envelope holds '
    'while the resource envelope does not. This is material rather than cosmetic, because bulk '
    'PII export is precisely the workload that processes GDPR-regulated personal data at volume, '
    'and it was measured on an environment already below the §4.5.c capacity floor — '
    'production headroom will be tighter, not looser. VULN-0044 must be fixed and STRESS-003 '
    're-run before bulk PII export is operated at production volume.',
    lead='STRESS-003 — Memory Threshold Breach (PASS with caveat): ')
# place the memory paragraph after the table
sla_tbl._tbl.addnext(p_mem._p)
log.append('§6: inserted STRESS-003 memory-threshold-breach paragraph (3.36 GB / 3.0 GB / 112% / VULN-0044)')

# ==================================================================
# 6. RISK MATRIX (Table 4)
# ==================================================================
t4 = [t for t in tables(d) if t.rows[0].cells[0].text.strip() == 'Risk Area'][0]

# R2 — AC-7 / VULN-0019 : tighten the mitigation cell
set_cell(t4.rows[2].cells[4],
    'Remediation targeted September 2, 2026 per the GDPR Privacy Audit Checklist — before '
    'the September 8 close and within the §4.4.b 30-day window. The CISO’s sign-off is '
    'separately conditioned on this fix; see §7.')
log.append('Risk matrix R2: replaced duplicated CISO paragraph with a remediation-focused mitigation')

# R12 — CISO conditional : the Risk Item cell held a whole paragraph
r12 = t4.rows[12]
assert 'VULN-0019' in cell_text(r12.cells[1]), cell_text(r12.cells[1])[:80]
set_cell(r12.cells[1], 'Priya Deshpande — CISO Conditional (VULN-0019)')
set_cell(r12.cells[4],
    'CISO approval is conditioned strictly on resolving VULN-0019 (PII export audit trail) before '
    'the closing date. The M&A Sign-off Tracker specifies no calendar date for the condition '
    '— only “before closing date.” The GDPR Privacy Audit Checklist separately '
    'records a September 2, 2026 target for AC-7/VULN-0019, which would precede the September 8 '
    'close. Confirm resolution and obtain unconditional sign-off before closing. See §7.')
log.append('Risk matrix R12: repaired malformed Risk Item cell and rewrote mitigation')

# R13 — Anita Flores : make consistent with §7's conclusion
r13 = t4.rows[13]
assert 'Anita Flores' in cell_text(r13.cells[1]), cell_text(r13.cells[1])[:80]
set_cell(r13.cells[4],
    'Condition is tied to VULN-0049, a Cosmetic UI defect targeted for September 15, 2026 — '
    '7 days after the September 8 hard close and therefore comfortably inside the 30-day '
    'post-close window the condition allows. Consistent with §7, this condition is '
    'satisfiable on the documented schedule; confirm completion post-close.')
log.append('Risk matrix R13: aligned Anita Flores row with §7 (Sep 15 satisfies the 30-day window)')

# New infrastructure row for the STRESS-003 memory breach
mem_row = clone_row(t4, 6)
set_cell(mem_row.cells[0], 'Infrastructure')
set_cell(mem_row.cells[1], 'STRESS-003 Bulk PII Export — Memory Threshold Breach / VULN-0044')
set_cell(mem_row.cells[2], 'High')
set_cell(mem_row.cells[3], 'Medium')
set_cell(mem_row.cells[4],
    'Throughput SLA met at all percentiles, but peak memory reached 3.36 GB against a 3.0 GB '
    'threshold (112%) and did not return to baseline. Corresponds to VULN-0044 (memory leak '
    'during sustained bulk PII export). Measured on a sub-80% environment, so production '
    'headroom is tighter. Fix VULN-0044 and re-run STRESS-003 before production-volume PII '
    'export. See §6.')
move_row_after(t4, mem_row, 6)
log.append('Risk matrix: inserted STRESS-003 / VULN-0044 infrastructure risk row')

# ==================================================================
# 7. FINAL RECOMMENDATION (§9)
# ==================================================================
p50 = find_para(d, 'GDPR SC-10 Not Met / VULN-0031 (Blocker)')
set_runs(p50, 'GDPR SC-10 Not Met / VULN-0031 (Blocker): ',
    'Policy §4.4.c requires an automatic NO-GO for a Not Met GDPR control absent a formal '
    'written CISO waiver. VULN-0031 causes SC-10 to be Not Met, and no waiver is documented. '
    'This is non-negotiable and sufficient on its own to halt the acquisition. Policy §5 '
    'fails independently: it requires every effective-Blocker vulnerability to be resolved prior '
    'to the Board Review, and VULN-0031’s remediation target of September 6, 2026 falls '
    'three days after the September 3, 2026 Board Review, so the Blocker will still be open when '
    'the Board convenes. Additionally, VULN-0031 is absent from the Known Quirks list in the '
    'Acquisition Announcement Draft — a material disclosure gap.')
log.append('§9 P50: added independent Policy §5 / Board Review failure')

p56 = find_para(d, 'Additionally, the following conditions must be satisfied')
p56.runs[0].text = (
    'Additionally, the following conditions must be satisfied before any revised GO '
    'consideration: VULN-0031 resolved (or a formal written CISO waiver obtained) in advance of '
    'the Board Review; VULN-0012 and VULN-0028 assigned to specific GDPR controls or formally '
    'scoped; Priya Deshpande’s CISO conditional cleared; Anita Flores’s QA conditional '
    'verified; VULN-0044 remediated and STRESS-003 re-run so that bulk PII export sits inside its '
    'memory threshold; the stress suite re-executed on an environment at or above the '
    '§4.5.c 80% floor, with VP Engineering’s P99 risk acceptance re-confirmed against '
    'those results; and VULN-0037 removed from the PR draft’s Known Quirks characterization.'
)
log.append('§9 P56: extended remaining-conditions list (VULN-0044, §4.5.c re-run, §4.5.b re-confirmation)')

d.save(P)

print('APPLIED EDITS')
for i, m in enumerate(log, 1):
    print(f'  {i:2d}. {m}')
print('\nSaved:', P)
