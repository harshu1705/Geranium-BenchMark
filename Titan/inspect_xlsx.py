import sys, io, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl

p = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'inputs', 'Cloud_Infrastructure_Stress_Test_Aug2026.xlsx')
wb = openpyxl.load_workbook(p)
for ws in wb.worksheets:
    print("=== SHEET:", ws.title, ws.dimensions)
    print("  freeze:", ws.freeze_panes, "| autofilter:", ws.auto_filter.ref)
    print("  col widths:", {k: v.width for k, v in ws.column_dimensions.items()})
    print("  sheetview tabColor:", ws.sheet_properties.tabColor)
    for row in ws.iter_rows(min_row=1, max_row=min(4, ws.max_row)):
        for c in row:
            if c.value is not None:
                print(f"   {c.coordinate}: font(b={c.font.bold},sz={c.font.size},name={c.font.name},color={c.font.color.rgb if c.font.color else None})"
                      f" fill={c.fill.fgColor.rgb if c.fill and c.fill.fgColor else None} patt={c.fill.patternType if c.fill else None}"
                      f" align(h={c.alignment.horizontal},wrap={c.alignment.wrap_text},v={c.alignment.vertical})"
                      f" border_bottom={c.border.bottom.style if c.border and c.border.bottom else None}")
        print("   ---")
    print("  merged:", ws.merged_cells.ranges)
    print("  row heights:", {k: v.height for k, v in ws.row_dimensions.items()})
