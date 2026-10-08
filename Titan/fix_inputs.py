import sys, io, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from copy import copy

BASE = os.path.dirname(os.path.abspath(__file__))
P = os.path.join(BASE, 'inputs', 'Cloud_Infrastructure_Stress_Test_Aug2026.xlsx')

wb = openpyxl.load_workbook(P)
ws = wb['Stress Test Results']

# ---------------------------------------------------------------
# 1. De-leak the STRESS-003 note.
#    The old note pre-stated the graded conclusion verbatim:
#      "throughput passes but memory usage hit 3.36GB vs 3.0GB
#       threshold (112%). See VULN-0044. asterisk = pass with caveat"
#    The raw numbers move to a new telemetry tab; the note now only
#    records the raw observation and points at where the counters live.
# ---------------------------------------------------------------
target = None
for r in range(2, ws.max_row + 1):
    if ws.cell(row=r, column=1).value == 'STRESS-003':
        target = r
        break
assert target is not None, 'STRESS-003 row not found'
print('STRESS-003 at row', target)
print('OLD NOTE:', ws.cell(row=target, column=10).value)

ws.cell(row=target, column=10).value = (
    'throughput clears the 500 rec/min bar at all three percentiles. '
    'resource counters for this run are logged on the Resource Utilization tab. '
    'asterisk = pass with caveat'
)
print('NEW NOTE:', ws.cell(row=target, column=10).value)

# ---------------------------------------------------------------
# 2. Add the raw "Resource Utilization" telemetry tab.
# ---------------------------------------------------------------
if 'Resource Utilization' in wb.sheetnames:
    del wb['Resource Utilization']

ru = wb.create_sheet('Resource Utilization')

hdr_font = Font(bold=True, size=11, name='Calibri', color='FFFFFF')
hdr_fill = PatternFill('solid', fgColor='4472C4')
hdr_align = Alignment(horizontal='center', wrap_text=True)
thin_bottom = Border(bottom=Side(style='thin'))
body_align = Alignment(wrap_text=True, vertical='top')

headers = [
    'Scenario ID', 'Scenario Name', 'Peak Memory (GB)', 'Memory Threshold (GB)',
    'Peak CPU (%)', 'CPU Threshold (%)', 'Measurement Window', 'Sample Interval', 'Notes',
]
widths = [14.0, 32.0, 18.0, 20.0, 14.0, 16.0, 20.0, 15.0, 60.0]

rows = [
    ('STRESS-001', 'Customer Record Search API',       1.42, 3.0, 61, 85, '60 min',  '10s', ''),
    ('STRESS-002', 'Executive Dashboard Load',         0.94, 3.0, 44, 85, '60 min',  '10s', ''),
    ('STRESS-003', 'Bulk PII Export',                  3.36, 3.0, 72, 85, '180 min', '10s',
     'peak recorded at minute 147 of the sustained export run; resident set did not return to '
     'pre-run baseline after the run completed'),
    ('STRESS-004', 'API Gateway — Concurrent Users', 2.10, 3.0, 68, 85, '45 min',  '10s', ''),
    ('STRESS-005', 'Database Query Latency',           1.88, 3.0, 57, 85, '60 min',  '10s', ''),
    ('STRESS-006', 'Report Generation',                2.34, 3.0, 63, 85, '60 min',  '10s', ''),
    ('STRESS-007', 'Login/SSO Flow',                   0.62, 3.0, 29, 85, '30 min',  '10s', ''),
    ('STRESS-008', 'Real-time Notification Delivery',  1.15, 3.0, 48, 85, '60 min',  '10s',
     'network-bound scenario — see Environment Specs note on test bandwidth'),
]

for i, h in enumerate(headers, start=1):
    c = ru.cell(row=1, column=i, value=h)
    c.font = hdr_font
    c.fill = hdr_fill
    c.alignment = hdr_align
    c.border = thin_bottom
    ru.column_dimensions[c.column_letter].width = widths[i - 1]

for ri, row in enumerate(rows, start=2):
    for ci, val in enumerate(row, start=1):
        c = ru.cell(row=ri, column=ci, value=(val if val != '' else None))
        c.alignment = body_align
        c.border = thin_bottom
        if ci in (3, 4):
            c.number_format = '0.00'

foot = ru.max_row + 2
ru.cell(row=foot, column=1, value='Note:').font = Font(bold=True, size=11, name='Calibri')
ru.cell(row=foot, column=2, value=(
    'Memory threshold is the per-service resident-set ceiling agreed with Cloud Ops for the '
    'staging tier. CPU threshold is sustained utilisation on the application tier. Counters '
    'captured by the standard telemetry agent; figures are raw, not normalised for '
    'environment scale.'
)).alignment = Alignment(wrap_text=True, vertical='top')
ru.merge_cells(start_row=foot, start_column=2, end_row=foot, end_column=9)

wb.save(P)
print('\nSaved:', P)

# ---------------------------------------------------------------
# Verify
# ---------------------------------------------------------------
wb2 = openpyxl.load_workbook(P, data_only=True)
print('\nSheets:', wb2.sheetnames)
for name in wb2.sheetnames:
    s = wb2[name]
    print(f'\n--- {name} ({s.dimensions}) ---')
    for row in s.iter_rows():
        vals = [('' if c.value is None else str(c.value)) for c in row]
        if any(v.strip() for v in vals):
            print(' | '.join(vals))
