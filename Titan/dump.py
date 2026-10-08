import sys, io, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl
from docx import Document

base = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'inputs')

def dump_xlsx(p):
    print("\n" + "#"*90)
    print("### XLSX:", os.path.basename(p))
    print("#"*90)
    wb = openpyxl.load_workbook(p, data_only=True)
    for ws in wb.worksheets:
        print(f"\n--- SHEET: {ws.title}  (dims {ws.dimensions}) ---")
        for row in ws.iter_rows():
            vals = []
            for c in row:
                v = c.value
                if v is None:
                    v = ""
                vals.append(str(v))
            if any(v.strip() for v in vals):
                print(" | ".join(vals))

def dump_docx(p):
    print("\n" + "#"*90)
    print("### DOCX:", os.path.basename(p))
    print("#"*90)
    d = Document(p)
    body = d.element.body
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    ti = 0
    for child in body.iterchildren():
        tag = child.tag.split('}')[-1]
        if tag == 'p':
            para = Paragraph(child, d)
            if para.text.strip():
                print(f"[{para.style.name}] {para.text}")
        elif tag == 'tbl':
            t = Table(child, d)
            print(f"\n=== TABLE {ti} ===")
            for r in t.rows:
                print(" | ".join(c.text.replace("\n", " / ") for c in r.cells))
            print(f"=== END TABLE {ti} ===\n")
            ti += 1

for f in sorted(os.listdir(base)):
    p = os.path.join(base, f)
    if f.endswith('.xlsx'):
        dump_xlsx(p)
    elif f.endswith('.docx'):
        dump_docx(p)
