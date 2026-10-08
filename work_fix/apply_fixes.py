import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
import os

BASE = r'c:\Users\user\OneDrive\Desktop\geranium project\work_fix'

###############################################################################
# FIX 1: Broker Email - restore stripped dollar amounts
###############################################################################
broker_path = os.path.join(BASE, 'inputs', 'Broker_Email_Correspondence.docx')
doc = Document(broker_path)

replacements_broker = {
    'They have explicitly requested a  limit across the board':
        'They have explicitly requested a $15,000,000 limit across the board',
    'We are competing against a ,000 quote from their incumbent carrier':
        'We are competing against a $45,000 quote from their incumbent carrier',
    'I really need our quote to come in under that ,000 mark':
        'I really need our quote to come in under that $45,000 mark',
    'Let me know as soon as possible if you can offer the full  limit on a primary basis':
        'Let me know as soon as possible if you can offer the full $15,000,000 limit on a primary basis',
}

for para in doc.paragraphs:
    full_text = para.text
    new_text = full_text
    for old, new in replacements_broker.items():
        new_text = new_text.replace(old, new)
    if new_text != full_text:
        # merge runs, replace text in first run, blank the rest
        if para.runs:
            para.runs[0].text = new_text
            for r in para.runs[1:]:
                r.text = ''
        print("Fixed broker paragraph:", new_text[:100])

doc.save(broker_path)
print("Saved broker email.\n")

###############################################################################
# FIX 2: Golden Memo - restore stripped dollar amounts + fix cyber-total wording
#         + flip the "misses target / uncompetitive" sentence
###############################################################################
golden_path = os.path.join(BASE, 'golden', 'Underwriting_Memo_NovaMart.docx')
doc = Document(golden_path)

replacements_golden = {
    'Annual Revenue: ,400,000':
        'Annual Revenue: $502,400,000',
    'Requested Limit: ,000,000\nRecommendation: DECLINE':
        'Requested Limit: $15,000,000\nRecommendation: DECLINE',

    'Base Premium Calculation: ,120':
        'Base Premium Calculation: $25,120',
    'The TTM revenue derived from the Financial Statements Summary is ,400,000. Using the standard retail base rate of .05 per ,000 in revenue, the base premium is calculated as follows: (,400,000 / 1000) * 0.05 = ,120. The base premium is established at ,120.':
        'The TTM revenue derived from the Financial Statements Summary is $502,400,000. Using the standard retail base rate of $0.05 per $1,000 in revenue, the base premium is calculated as follows: ($502,400,000 / 1,000) * 0.05 = $25,120. The base premium is established at $25,120.',

    'Loss Run Surcharge Multiplier: 1.20x\nRationale: The Loss Run History reveals a 2025 cyber ransomware claim (CLM-25-1102) totaling ,000 paid, alongside three other in-window non-cyber claims (CLM-24-8891 GL, CLM-26-0045 Property, CLM-26-0891 Auto, plus CLM-23-4412 Cyber in the 3-year window). The total all-lines incurred losses over the past 3 years is ,500. Section 4 of the Underwriting Guidelines states that claims totaling >= ,000 in the last 3 years require a 1.20x penalty multiplier. Because the all-lines total (,500) exceeds the ,000 threshold, it warrants the maximum loss penalty.':
        'Loss Run Surcharge Multiplier: 1.20x\nRationale: The Loss Run History contains two cyber-only claims within the last 3 years: CLM-23-4412 (a January 2023 encrypted-laptop loss, $7,500 paid) and CLM-25-1102 (the 2025 ransomware event on the legacy ValueStore server, $650,000 paid). Whether the January 2023 claim falls inside the exact 3-year lookback window is a boundary judgment call: excluding it yields a cyber-only 3-year total of $650,000, while including it yields $657,500. Under Section 4 of the Underwriting Guidelines, both figures fall within the "Severe claims ($500k or more in the last 3 years)" band, so either interpretation results in the same 1.20x Loss Run Surcharge Multiplier. The three other in-window claims (CLM-24-8891 GL, CLM-26-0045 Property, CLM-26-0891 Auto) are non-cyber lines noted for file completeness but do not factor into this cyber-specific surcharge calculation.',

    'The Ransomware Supplemental App confirms that Multi-Factor Authentication (MFA) is missing for Third-Party/Contractor Remote Access.':
        'The Ransomware Supplemental App confirms that Multi-Factor Authentication (MFA) is missing for Third-Party/Contractor Remote Access.',

    'Final Calculated Premium: ,680 (Not Quoted due to Decline)\nRationale: Base (,120) * Loss Mult (1.20) * Security Mult (1.25) = ,680. The hypothetical premium is ,680. This significantly misses the broker\u2019s target of ,000. Even if the EDR gap were resolved, the account is uncompetitive and would face heavy pricing resistance from the insured.':
        'Final Calculated Premium: $37,680 (Not Quoted due to Decline)\nRationale: Base ($25,120) * Loss Mult (1.20) * Security Mult (1.25) = $37,680. The hypothetical premium is $37,680, which is $7,320 below the broker\u2019s target premium of $45,000 per the Broker Email Correspondence. On price alone this account would be competitive and would satisfy the broker\u2019s pricing objective; however, this is immaterial to the recommendation because the mandatory decline conditions below prevent the policy from being quoted at all.',

    'Eligible Limit: ,000,000 (Moot due to Decline)\nRationale: The missing Contractor MFA caps the maximum eligible limit at ,000,000 per Section 5 of the guidelines. The requested limit of ,000,000 is therefore ineligible, even setting aside the EDR decline trigger.':
        'Eligible Limit: $10,000,000 (Moot due to Decline)\nRationale: The missing Contractor MFA caps the maximum eligible limit at $10,000,000 per Section 5 of the guidelines. The requested limit of $15,000,000 is therefore ineligible, even setting aside the EDR decline trigger.',

    'Authority Required: Senior Underwriter (Moot due to Decline)\nRationale: The requested limit of ,000,000 requires Senior Underwriter authority per Section 7. While the actual eligible limit is capped at ,000,000, this still requires Senior Underwriter authority because it exceeds the ,000,000 capacity of an Underwriter I. The file is simply declined by the underwriter due to the EDR gap, and that a VP referral is NOT required for a single-condition decline. Furthermore, per Section 8 of the Underwriting Guidelines, because the risk triggers multiple simultaneous decline conditions (missing EDR and an unresolved critical CVE on the perimeter), immediate file closure is required with no VP referral permitted.':
        'Authority Required: Senior Underwriter (Moot due to Decline)\nRationale: The requested limit of $15,000,000 requires Senior Underwriter authority per Section 7. While the actual eligible limit is capped at $10,000,000, this still requires Senior Underwriter authority because it exceeds the $5,000,000 capacity of an Underwriter I. Regardless of authority level, per Section 8 of the Underwriting Guidelines, because the risk triggers multiple simultaneous decline conditions (missing EDR and an unresolved critical CVE on the perimeter), immediate file closure is required with no VP referral permitted.',
}

changed_count = 0
for para in doc.paragraphs:
    full_text = para.text
    new_text = full_text
    for old, new in replacements_golden.items():
        new_text = new_text.replace(old, new)
    if new_text != full_text:
        if para.runs:
            para.runs[0].text = new_text
            for r in para.runs[1:]:
                r.text = ''
        changed_count += 1
        print("Fixed golden paragraph:", new_text[:100])

# Fix Risk Matrix table
table_replacements = {
    ',500 total losses in 3 years triggers 1.20x penalty multiplier.':
        '$650,000\u2013$657,500 cyber-only losses in the last 3 years (depending on whether the January 2023 claim falls within the exact lookback window) triggers the maximum 1.20x penalty multiplier.',
    'Missing EDR on 2% of endpoints triggers hard decline. Missing MFA on Contractor Remote Access triggers 1.25x penalty and caps limit at ,000,000.':
        'Missing EDR on 2% of endpoints triggers hard decline. Missing MFA on Contractor Remote Access triggers 1.25x penalty and caps limit at $10,000,000.',
    'Final calculated premium is ,680, which exceeds the broker target premium of ,000.':
        'Final calculated premium is $37,680, which is $7,320 below the broker target premium of $45,000\u2014meeting the broker\u2019s pricing objective, though moot given the decline.',
}

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for old, new in table_replacements.items():
                if old in cell.text:
                    for p in cell.paragraphs:
                        if old in p.text:
                            t = p.text.replace(old, new)
                            if p.runs:
                                p.runs[0].text = t
                                for r in p.runs[1:]:
                                    r.text = ''
                            print("Fixed table cell ->", t[:100])

doc.save(golden_path)
print(f"\nSaved golden memo. {changed_count} paragraphs changed.")
