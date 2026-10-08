# -*- coding: utf-8 -*-
import csv, json, datetime
from fractions import Fraction as F
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

with open('work/grades.json') as f:
    G = json.load(f)
with open('work/costing.json') as f:
    C = json.load(f)

bands = G['bands']
grade_min = {int(k):int(v) for k,v in G['grade_min'].items()}
grade_mid = {int(k):int(v) for k,v in G['grade_mid'].items()}
grade_max = {int(k):int(v) for k,v in G['grade_max'].items()}
grade_min_pref = {int(k):int(v) for k,v in G['grade_min_pref'].items()}
usable_count = {int(k):v for k,v in G['usable_count'].items()}
mechanism = {int(k):v for k,v in G['mechanism'].items()}
zone_table = G['zone_table']
grade_of = G['grade_of']
points_of = G['points_of']
titles = G['titles']
excluded = G['excluded']
step_repair_detail = {int(k):v for k,v in G['step_repair_detail'].items()}
interp_from = {int(k):v for k,v in G['interp_from'].items()}
n_grades = len(bands)

ZONES = ['Zone 1','Zone 2','Zone 3']
ZONE_STATES = {'Zone 1': ['Ohio','Indiana','Missouri'],
               'Zone 2': ['Colorado','Illinois','Minnesota'],
               'Zone 3': ['California','Washington','New York']}
ZONE_FACTOR = {'Zone 1':0.92,'Zone 2':1.00,'Zone 3':1.12}

# ---------- styling helpers ----------
NAVY = "1F3864"
LIGHT = "DCE6F1"
AMBER = "FFF2CC"
RED = "FCE4E4"
GREEN = "E2EFDA"
thin = Side(style='thin', color="B7B7B7")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def header_row(ws, row, values, fill=NAVY, font_color="FFFFFF", widths=None):
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = Font(bold=True, color=font_color, size=10)
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = border

def title_cell(ws, row, col, text, size=14, color=NAVY):
    c = ws.cell(row=row, column=col, value=text)
    c.font = Font(bold=True, size=size, color=color)

def note_cell(ws, row, col, text, italic=True, size=9, color="595959"):
    c = ws.cell(row=row, column=col, value=text)
    c.font = Font(italic=italic, size=size, color=color)

def money(ws, row, col, value, fmt='#,##0'):
    c = ws.cell(row=row, column=col, value=value)
    c.number_format = fmt
    c.border = border
    return c

def plain(ws, row, col, value, bold=False, align='left', wrap=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(bold=bold, size=10)
    c.alignment = Alignment(horizontal=align, vertical='center', wrap_text=wrap)
    c.border = border
    return c

wb = openpyxl.Workbook()
wb.remove(wb.active)

# =========================================================
# SHEET 1: Cover & Summary
# =========================================================
ws = wb.create_sheet("Cover & Summary")
ws.sheet_view.showGridLines = False
title_cell(ws, 2, 2, "Millgate Mutual - 2027 Salary Structure Specification", 16)
note_cell(ws, 3, 2, "Prepared by Total Rewards. Structure effective 1 January 2027. Adjustment payments effective 1 February 2027 (Year 1) and 1 February 2028 (Year 2, where phased).")
note_cell(ws, 4, 2, "Source files: Salary_Structure_Design_Standard_2027.docx; Total_Rewards_Budget_Memo.docx; Market_Survey_Match_Extract.xlsx; Job_Evaluation_Point_Schedule.csv; Incumbent_Roster.csv; State_Wage_Floor_Schedule_2027.csv.")

rows = [
    ("Number of grades built from the point schedule", f"{n_grades}"),
    ("Point range on the schedule (66 evaluated jobs)", f"{bands[0][0]+ (min(int(v) for v in points_of.values())-bands[0][0])}-{max(int(v) for v in points_of.values())} pts (floor {bands[0][0]}, ceiling {bands[-1][1]})"),
    ("Grade 1 band", f"{bands[0][0]}-{bands[0][1]} points"),
    (f"Grade {n_grades} band", f"{bands[-1][0]}-{bands[-1][1]} points"),
    ("Survey matches in extract", "62, of which 3 are set aside (see Benchmark Eligibility)"),
    ("Usable benchmark jobs feeding the market build", "33"),
    ("Incumbents on the roster", "150"),
    ("Excluded from cost - employment ends before 1 Feb 2027", "3"),
    ("Excluded from grading - job code carries no evaluation points", "2 (flagged on the Ungradable Incumbents tab)"),
    ("Costed population (in employment 1 Feb 2027 and gradable)", "145"),
    ("Incumbents whose salary sits below the new zone minimum", "57"),
    ("Cost of moving every affected incumbent to the new minimum", "$391,600"),
    ("Board-approved implementation budget", "$355,000"),
    ("Cost exceeds budget by", "$36,600 - phase-in applies per the budget memo"),
    ("Paid 1 February 2027", "$208,600"),
    ("Paid 1 February 2028", "$183,000"),
]
r = 6
header_row(ws, r, ["Summary item", "Value"])
for i,(k,v) in enumerate(rows, start=1):
    plain(ws, r+i, 1, k)
    plain(ws, r+i, 2, v, align='left', wrap=True)
ws.column_dimensions['A'].width = 52
ws.column_dimensions['B'].width = 58
for i in range(len(rows)+1):
    ws.row_dimensions[r+i].height = 16

r2 = r + len(rows) + 3
title_cell(ws, r2, 2, "Contents", 12)
contents = ["Grade & Zone Ranges (loader-ready, no blanks)",
            "Benchmark Eligibility (all 62 survey matches, 3 set aside with reason)",
            "Market Build & Mechanism (working behind every grade)",
            "Ungradable Incumbents (roster entries that cannot be placed)",
            "Implementation Cost (population, shortfall, phase-in)"]
for i,c in enumerate(contents, start=1):
    plain(ws, r2+i, 2, "• " + c)

# =========================================================
# SHEET 2: Grade & Zone Ranges
# =========================================================
ws = wb.create_sheet("Grade & Zone Ranges")
ws.sheet_view.showGridLines = False
title_cell(ws, 1, 1, "Millgate Mutual 2027 Salary Structure - Grade & Zone Ranges", 14)
note_cell(ws, 2, 1, "Effective 1 January 2027. Every grade carries a minimum, midpoint and maximum in all three zones; no cell is left blank.")

headers = ["Grade","Point band","Setting mechanism",
           "Zone 1 Min","Zone 1 Mid","Zone 1 Max",
           "Zone 2 Min","Zone 2 Mid","Zone 2 Max",
           "Zone 3 Min","Zone 3 Mid","Zone 3 Max","Statutory floor note"]
hr = 4
header_row(ws, hr, headers)
for i,g in enumerate(range(1,n_grades+1), start=1):
    row = hr+i
    plain(ws, row, 1, g, bold=True, align='center')
    plain(ws, row, 2, f"{bands[g-1][0]}-{bands[g-1][1]}", align='center')
    plain(ws, row, 3, mechanism[g], align='center')
    col = 4
    floor_notes = []
    for z in ZONES:
        d = zone_table[f"{g}|{z}"]
        money(ws, row, col, d['min']); col+=1
        money(ws, row, col, d['mid']); col+=1
        money(ws, row, col, d['max']); col+=1
        if d['floor_applied']:
            ft = "exempt salary threshold" if d['floor_type']=='salary' else "annualized hourly minimum"
            floor_notes.append(f"{z} minimum lifted to the {ft} (${d['floor_val']:,.0f}) - highest applicable state in the zone")
    plain(ws, row, 13, "; ".join(floor_notes) if floor_notes else "-", wrap=True)
for col in range(1,14):
    ws.column_dimensions[get_column_letter(col)].width = 13
ws.column_dimensions['A'].width = 8
ws.column_dimensions['B'].width = 12
ws.column_dimensions['C'].width = 16
ws.column_dimensions['M'].width = 46
for i in range(n_grades):
    ws.row_dimensions[hr+1+i].height = 26

note_cell(ws, hr+n_grades+3, 1, "Zone factors: Zone 1 = 0.92 (Ohio, Indiana, Missouri); Zone 2 = 1.00 (Colorado, Illinois, Minnesota); Zone 3 = 1.12 (California, Washington, New York).")
note_cell(ws, hr+n_grades+4, 1, "Spreads: grades 1-3 = 40%; grades 4-7 = 45%; grades 8-11 = 50%. Minimum = midpoint / (1 + spread/2). Maximum = minimum x (1 + spread).")

# =========================================================
# SHEET 3: Benchmark Eligibility
# =========================================================
ws = wb.create_sheet("Benchmark Eligibility")
ws.sheet_view.showGridLines = False
title_cell(ws, 1, 1, "Market Survey Match Extract - Benchmark Eligibility Screen", 14)
note_cell(ws, 2, 1, "A match is usable only where the source is effective within 24 months of 1 Jan 2027, the scope match is Exact or Good, and the cut carries at least 5 incumbents. All 62 matches in the extract are screened below; 3 are set aside.")

hdr = ["Job code","Benchmark title","Eval. points","Source","Source effective","Reported 50th pct.","Scope match","Incumbents in cut","Usable?","Reason set aside"]
hr = 4
header_row(ws, hr, hdr)

survey_rows = []
with open('inputs_extract/market_survey_match_detail.csv') as f:
    for row in csv.DictReader(f):
        survey_rows.append(row)

set_aside_reason = {
    'AC-610': "Source (Hollis Regional Property and Casualty Study, effective 2024-10-01) is more than 24 months before the 1 Jan 2027 structure effective date - stale.",
    'CL-210': "Cut carries only 3 incumbents, below the 5-incumbent minimum - too thin to use.",
    'CL-960': "Scope match recorded as Partial, not Exact or Good - a scope we would not defend.",
}

ridx = hr
for row in survey_rows:
    ridx += 1
    usable = row['job_code'] not in set_aside_reason
    plain(ws, ridx, 1, row['job_code'], align='center')
    plain(ws, ridx, 2, row['benchmark_title'])
    plain(ws, ridx, 3, int(row['evaluation_points']), align='center')
    plain(ws, ridx, 4, row['source'])
    plain(ws, ridx, 5, row['source_effective'], align='center')
    money(ws, ridx, 6, int(row['reported_50th_percentile']))
    plain(ws, ridx, 7, row['scope_match'], align='center')
    plain(ws, ridx, 8, int(row['incumbents_in_cut']), align='center')
    cell = plain(ws, ridx, 9, "Yes" if usable else "No", bold=True, align='center')
    if not usable:
        for cc in range(1,11):
            ws.cell(row=ridx, column=cc).fill = PatternFill("solid", fgColor=RED)
        plain(ws, ridx, 10, set_aside_reason[row['job_code']], wrap=True)
    else:
        plain(ws, ridx, 10, "-", align='center')

widths = [10,26,10,30,13,13,10,12,8,50]
for i,w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

note_cell(ws, ridx+2, 1, "62 matches in the extract, spanning 36 distinct benchmark jobs. 3 job codes (AC-610, CL-210, CL-960) are set aside and take no part in any calculation; 33 benchmark jobs remain usable.")

# =========================================================
# SHEET 4: Market Build & Mechanism
# =========================================================
ws = wb.create_sheet("Market Build & Mechanism")
ws.sheet_view.showGridLines = False
title_cell(ws, 1, 1, "Working Behind Each Grade's Minimum, Midpoint and Maximum", 14)
note_cell(ws, 2, 1, "Aged rate = reported 50th percentile x (1 + 3.2% x months since source effective date / 12). Blend = Cranmere 60% / Delaney 40% where both are usable; full weight where only one source is usable.")

# Precompute per-job aged/blended values for display, matching build5 methodology
def months_between(d1,d2):
    return (d2.year-d1.year)*12+(d2.month-d1.month)
EFFECTIVE = datetime.date(2027,1,1)
def parse(d):
    y,m,dd = d.split('-'); return datetime.date(int(y),int(m),int(dd))

WEIGHTS = {'Cranmere Carrier Compensation Survey':F(6,10), 'Delaney National Insurance Survey':F(4,10)}
usable_by_job = {}
for row in survey_rows:
    if row['job_code'] in set_aside_reason: continue
    age_m = months_between(parse(row['source_effective']), EFFECTIVE)
    aged = F(int(row['reported_50th_percentile'])) * (1+F(32,1000)*F(age_m,12))
    usable_by_job.setdefault(row['job_code'], []).append((row['source'], aged))

blended = {}
for code, entries in usable_by_job.items():
    if len(entries)==1:
        blended[code] = entries[0][1]
    else:
        num = sum(a*WEIGHTS[s] for s,a in entries)
        blended[code] = num  # weights sum to 1

hr = 4
hdr = ["Grade","Usable benchmark jobs (points, blended aged rate $)","Raw median ($)","Grade midpoint (rounded, $)","Minimum derivation","Grade minimum ($)","Maximum derivation","Grade maximum ($)","Mechanism"]
header_row(ws, hr, hdr)
ridx = hr
grade_jobs = {g: [] for g in range(1,n_grades+1)}
for code, rate in blended.items():
    g = grade_of[code]
    grade_jobs[g].append((code, points_of[code], rate))

for g in range(1,n_grades+1):
    ridx += 1
    jobs = sorted(grade_jobs[g], key=lambda x: x[2])
    spread = 0.40 if g<=3 else (0.45 if g<=7 else 0.50)
    if jobs:
        jobtext = "; ".join(f"{c} ({p} pts): ${float(v):,.2f}" for c,p,v in jobs)
    else:
        b,a = interp_from[g]
        jobtext = f"No usable benchmark in band. Interpolated: (Grade {b} midpoint ${grade_mid[b]:,} + Grade {a} midpoint ${grade_mid[a]:,}) / 2"
    plain(ws, ridx, 1, g, bold=True, align='center')
    plain(ws, ridx, 2, jobtext, wrap=True)
    if jobs:
        n = len(jobs)
        vals = [v for _,_,v in jobs]
        med = vals[n//2] if n%2==1 else (vals[n//2-1]+vals[n//2])/2
        plain(ws, ridx, 3, f"${float(med):,.2f}", align='right')
    else:
        plain(ws, ridx, 3, "n/a - interpolated", align='center')
    if g in step_repair_detail:
        direction, cur, prev = step_repair_detail[g]
        plain(ws, ridx, 4, f"${grade_mid[g]:,} - Step Repair: market/raw midpoint (${cur:,}) was a {(cur/prev-1)*100:.2f}% step from Grade {g-1} (${prev:,}), outside the 8-14% band; moved to the nearest permitted boundary and rounded inward.", wrap=True)
    else:
        plain(ws, ridx, 4, f"${grade_mid[g]:,}", align='center')
    plain(ws, ridx, 5, f"${grade_mid[g]:,} / (1 + {spread/2:.3f}) = ${grade_min_pref[g]:,}", wrap=True)
    if mechanism[g]=='Continuity Floor':
        plain(ws, ridx, 6, f"${grade_min[g]:,} - Continuity Floor: range geometry gives ${grade_min_pref[g]:,}, below the ${grade_min[g]:,} minimum in force for this grade number in the prior structure; the figure in force governs.", wrap=True)
    else:
        plain(ws, ridx, 6, f"${grade_min[g]:,}", align='center')
    plain(ws, ridx, 7, f"${grade_min[g]:,} x (1 + {spread:.2f}) = ${grade_max[g]:,}", wrap=True)
    plain(ws, ridx, 8, f"${grade_max[g]:,}", align='center')
    plain(ws, ridx, 9, mechanism[g], align='center', bold=True)

widths = [7,55,12,32,26,12,26,12,15]
for i,w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
for i in range(n_grades):
    ws.row_dimensions[hr+1+i].height = 46

# =========================================================
# SHEET 5: Ungradable Incumbents
# =========================================================
ws = wb.create_sheet("Ungradable Incumbents")
ws.sheet_view.showGridLines = False
title_cell(ws, 1, 1, "Incumbents That Cannot Be Placed on the 2027 Structure", 14)
note_cell(ws, 2, 1, "These job codes carry no evaluation point total on the Job Evaluation Point Schedule and therefore cannot be assigned a grade under Band Construction. No grade has been guessed. Pilar should chase the missing field before the 19 December load window closes.")

hdr = ["Employee ID","Job code","Job title (roster)","Work state","Current base salary","Missing field","Action needed"]
hr=4
header_row(ws, hr, hdr)
roster_rows = []
with open('inputs_extract/Incumbent_Roster.csv') as f:
    for row in csv.DictReader(f):
        roster_rows.append(row)
ridx = hr
for eid in C['ungradable']:
    r = next(x for x in roster_rows if x['employee_id']==eid)
    ridx += 1
    plain(ws, ridx, 1, r['employee_id'], bold=True, align='center')
    plain(ws, ridx, 2, r['job_code'], align='center')
    plain(ws, ridx, 3, r['job_title'])
    plain(ws, ridx, 4, r['work_state'], align='center')
    money(ws, ridx, 5, int(r['current_base_salary']))
    plain(ws, ridx, 6, "evaluation_points - job code absent from Job_Evaluation_Point_Schedule.csv", wrap=True)
    plain(ws, ridx, 7, "Committee to evaluate and publish points for this job code before load; no grade assigned in the interim.", wrap=True)
widths=[12,10,26,12,16,40,44]
for i,w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
for i in range(len(C['ungradable'])):
    ws.row_dimensions[hr+1+i].height = 30

# =========================================================
# SHEET 6: Implementation Cost
# =========================================================
ws = wb.create_sheet("Implementation Cost")
ws.sheet_view.showGridLines = False
title_cell(ws, 1, 1, "Implementation Cost and Phase-In", 14)

pop = [
    ("Incumbents on the roster", 150),
    ("Less: employment ends before 1 Feb 2027 (not costed)", -3),
    ("In employment on 1 Feb 2027", 147),
    ("Less: job code carries no evaluation points (ungradable, not costed)", -2),
    ("Costed population", 145),
    ("Of which: salary already at or above the new zone minimum (no adjustment)", 88),
    ("Of which: salary below the new zone minimum (affected)", 57),
]
r=3
header_row(ws, r, ["Population","Count"])
for i,(k,v) in enumerate(pop, start=1):
    plain(ws, r+i, 1, k)
    plain(ws, r+i, 2, v, align='center')
ws.column_dimensions['A'].width = 62

r2 = r+len(pop)+2
header_row(ws, r2, ["Budget test","Amount"])
budget_rows = [
    ("Cost of moving every affected incumbent to the new zone minimum", 391600),
    ("Board-approved implementation budget", 355000),
    ("Amount by which cost exceeds budget", 36600),
]
for i,(k,v) in enumerate(budget_rows, start=1):
    plain(ws, r2+i, 1, k)
    money(ws, r2+i, 2, v)
note_cell(ws, r2+len(budget_rows)+1, 1, "Cost exceeds the approved budget, so the phase-in in the Total Rewards Budget Memo applies.")

r3 = r2+len(budget_rows)+3
header_row(ws, r3, ["Phase-in rule","Result"])
phase_rows = [
    ("On 1 Feb 2027: greater of half the shortfall, and the amount needed to reach the statutory floor applicable to the employee's own work state", "See incumbent detail below"),
    ("Remainder paid 1 Feb 2028", "See incumbent detail below"),
    ("Total paid 1 February 2027", 208600),
    ("Total paid 1 February 2028", 183000),
]
for i,(k,v) in enumerate(phase_rows, start=1):
    plain(ws, r3+i, 1, k, wrap=True)
    if isinstance(v,int):
        money(ws, r3+i, 2, v)
    else:
        plain(ws, r3+i, 2, v, align='center')

r4 = r3+len(phase_rows)+3
title_cell(ws, r4, 1, "Affected Incumbent Detail (57 employees) - by Employee ID, not name", 12)
hdr = ["Employee ID","Job code","Grade","Work state","Zone","Current base salary","New zone minimum","Shortfall","Statutory floor applicable to employee (own state)","Floor top-up","Half of shortfall","Paid 1 Feb 2027","Paid 1 Feb 2028"]
hr = r4+2
header_row(ws, hr, hdr)
ridx = hr
for row in sorted(C['rows2027'], key=lambda x: x[0]):
    (eid,name,jobcode,state,zone,g,salary,zmin,shortfall,floorval,floortype,floortopup,halfr,amt27,amt28) = row
    ridx += 1
    plain(ws, ridx, 1, eid, align='center')
    plain(ws, ridx, 2, jobcode, align='center')
    plain(ws, ridx, 3, g, align='center')
    plain(ws, ridx, 4, state, align='center')
    plain(ws, ridx, 5, zone, align='center')
    money(ws, ridx, 6, salary)
    money(ws, ridx, 7, zmin)
    money(ws, ridx, 8, shortfall)
    plain(ws, ridx, 9, f"${floorval:,.0f} ({floortype})", align='center')
    money(ws, ridx, 10, floortopup)
    money(ws, ridx, 11, halfr)
    c27 = money(ws, ridx, 12, amt27)
    if floortopup > halfr:
        c27.fill = PatternFill("solid", fgColor=AMBER)
    money(ws, ridx, 13, amt28)

totrow = ridx+1
plain(ws, totrow, 5, "TOTAL", bold=True, align='right')
money(ws, totrow, 8, sum(r[8] for r in C['rows2027']))
money(ws, totrow, 12, sum(r[13] for r in C['rows2027']))
money(ws, totrow, 13, sum(r[14] for r in C['rows2027']))
for c in range(1,14):
    ws.cell(row=totrow, column=c).font = Font(bold=True)

widths = [11,10,7,11,9,15,14,12,30,11,13,13,13]
for i,w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

note_cell(ws, totrow+2, 1, "Amber cells: the employee's own state statutory floor top-up exceeds half the shortfall, so the floor (not half-the-shortfall) sets the 1 Feb 2027 payment. The statutory floor used here is the employee's own work-state threshold, not the (possibly higher) threshold used to set the zone minimum.")
note_cell(ws, totrow+3, 1, "Half-of-shortfall and floor top-up are each rounded to the nearest hundred dollars, halves upward, consistent with the design standard's rounding rule.")

wb.save('Millgate_2027_Salary_Structure_Specification.xlsx')
print("Saved workbook.")
print("Total affected rows written:", ridx-hr)
