import openpyxl
from datetime import datetime

wb = openpyxl.load_workbook('Crestline_AR_Ledger_Hartsfield_History.xlsx')

DATE_FMT = 'yyyy-mm-dd'
CUR = '#,##0'

ws = wb['Crestline AR Ledger']

for r in range(7, 20):
    inv_cell = ws[f'B{r}']
    due_cell = ws[f'D{r}']
    paid_cell = ws[f'E{r}']
    days_cell = ws[f'F{r}']

    # Convert Invoice Date string -> real date, format it
    inv_date = datetime.strptime(inv_cell.value, '%Y-%m-%d')
    inv_cell.value = inv_date
    inv_cell.number_format = DATE_FMT

    # Due Date becomes a live formula off Invoice Date (Net 30 terms)
    due_cell.value = f'=B{r}+30'
    due_cell.number_format = DATE_FMT

    # Date Paid: real date where paid, else keep 'OUTSTANDING' text
    if paid_cell.value != 'OUTSTANDING':
        paid_date = datetime.strptime(paid_cell.value, '%Y-%m-%d')
        paid_cell.value = paid_date
        paid_cell.number_format = DATE_FMT
        days_cell.value = f'=E{r}-B{r}'
    else:
        days_cell.value = '=IF(E{0}="OUTSTANDING","N/A",E{0}-B{0})'.format(r)

    days_cell.number_format = '0'
    ws[f'C{r}'].number_format = CUR

ws2 = wb['Account Summary']
ws2['B3'].number_format = CUR
ws2['B4'].value = "='Crestline AR Ledger'!C17+'Crestline AR Ledger'!C18+'Crestline AR Ledger'!C19"
ws2['B4'].number_format = CUR
ws2['B5'].value = '=B3-B4'
ws2['B5'].number_format = CUR
for r in range(8, 13):
    ws2[f'B{r}'].number_format = CUR
ws2['B13'].value = '=SUM(B8:B12)'
ws2['B13'].number_format = CUR

wb.save('Crestline_AR_Ledger_Hartsfield_History.xlsx')
print('saved')
