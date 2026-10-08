import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from datetime import datetime, timedelta

def create_financials():
    wb = openpyxl.Workbook()
    
    # Sheet 1: Income Statement
    ws_is = wb.active
    ws_is.title = "Income Statement"
    
    # Title
    ws_is.append(['Hartsfield Polymer Technologies, LLC', ''])
    ws_is.append(['Income Statement (Unaudited)', ''])
    ws_is.append(['For the year ended Dec 31, 2025 (in USD)', ''])
    ws_is.append(['', ''])
    
    # Data
    is_data = [
        ['Sales Revenue', 10435200],
        ['Cost of Goods Sold (COGS)', 7215000],
        ['Gross Profit', 3220200],
        ['Selling, General & Admin (SG&A)', 2195000],
        ['Depreciation & Amortization', 219750],
        ['Earnings Before Interest and Taxes (EBIT)', 805450],
        ['Interest Expense', 110500],
        ['Earnings Before Taxes (EBT)', 694950],
        ['Tax Expense', 198000],
        ['Net Income', 496950]
    ]
    for row in is_data:
        ws_is.append(row)
        
    # Formatting
    for row in ws_is.iter_rows(min_row=1, max_row=3, max_col=1):
        for cell in row:
            cell.font = Font(bold=True)
            
    # Sheet 2: Balance Sheet
    ws_bs = wb.create_sheet(title="Balance Sheet")
    
    # Title
    ws_bs.append(['Hartsfield Polymer Technologies, LLC', '', ''])
    ws_bs.append(['Balance Sheet (Unaudited)', '', ''])
    ws_bs.append(['As of Dec 31 (in USD)', '2025', '2024'])
    ws_bs.append(['', '', ''])
    
    # Data (added 2024 comparatives)
    bs_data = [
        ['ASSETS', '', ''],
        ['Cash and Cash Equivalents', 485200, 431500],
        ['Accounts Receivable', 1515000, 1420100],
        ['Inventory', 1935000, 1851000],
        ['Prepaid Expenses', 180000, 145000],
        ['Total Current Assets', 4115200, 3847600],
        ['Property, Plant & Equipment (Net)', 5235300, 5812000],
        ['Intangibles & Goodwill', 832000, 515000],
        ['Total Assets', 10182500, 10174600],
        ['', '', ''],
        ['LIABILITIES', '', ''],
        ['Accounts Payable', 1185000, 1095000],
        ['Short-Term Debt', 1525500, 1750000],
        ['Accrued Liabilities', 474000, 482000],
        ['Total Current Liabilities', 3184500, 3327000],
        ['Long-Term Debt', 4061100, 3755000],
        ['Total Liabilities', 7245600, 7082000],
        ['', '', ''],
        ['EQUITY', '', ''],
        ['Common Stock', 2000000, 2000000],
        ['Retained Earnings', 936900, 1092600],
        ['Total Equity (Book Value)', 2936900, 3092600],
        ['Total Liabilities & Equity', 10182500, 10174600]
    ]
    for row in bs_data:
        ws_bs.append(row)
        
    for row in ws_bs.iter_rows(min_row=1, max_row=3):
        for cell in row:
            cell.font = Font(bold=True)
            
    # Sheet 3: H1 2026 Interim
    ws_interim = wb.create_sheet(title="H1 2026 Interim IS")
    ws_interim.append(['Hartsfield Polymer Technologies, LLC', ''])
    ws_interim.append(['Interim Income Statement (Unaudited)', ''])
    ws_interim.append(['For the 6 months ended Jun 30, 2026 (in USD)', ''])
    ws_interim.append(['', ''])
    
    interim_data = [
        ['Sales Revenue', 5865200],
        ['Cost of Goods Sold (COGS)', 4452100],
        ['Gross Profit', 1413100],
        ['Selling, General & Admin (SG&A)', 1162500],
        ['Depreciation & Amortization', 105000],
        ['Earnings Before Interest and Taxes (EBIT)', 145600],
        ['Interest Expense', 135000],
        ['Earnings Before Taxes (EBT)', 10600],
        ['Tax Expense', 2226],
        ['Net Income', 8374]
    ]
    for row in interim_data:
        ws_interim.append(row)
        
    for row in ws_interim.iter_rows(min_row=1, max_row=3):
        for cell in row:
            cell.font = Font(bold=True)
            
    # Adjust column widths
    for ws in [ws_is, ws_bs, ws_interim]:
        ws.column_dimensions['A'].width = 40
        ws.column_dimensions['B'].width = 15
        if ws == ws_bs:
            ws.column_dimensions['C'].width = 15
            
    wb.save('Hartsfield_Polymer_Unaudited_Financials.xlsx')

def create_ar_ledger():
    wb = openpyxl.Workbook()
    
    # Sheet 1: Ledger
    ws = wb.active
    ws.title = "Crestline AR Ledger"
    
    ws.append(['Crestline Chemical Supply, Inc. - Internal AR Ledger', '', '', '', '', ''])
    ws.append(['Customer: Hartsfield Polymer Technologies, LLC (Acct # 88392)', '', '', '', '', ''])
    ws.append(['Current Terms: Net 30', '', '', '', '', ''])
    ws.append(['As of Date: August 20, 2026', '', '', '', '', ''])
    ws.append(['', '', '', '', '', ''])
    
    headers = ['Invoice #', 'Invoice Date', 'Amount (USD)', 'Due Date', 'Date Paid', 'Days to Pay (From Inv Date)']
    ws.append(headers)
    
    # Formatting headers
    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
    for cell in ws[6]:
        cell.font = header_font
        cell.fill = header_fill
        
    # Invoice Data (includes edge cases for strict temporal reasoning)
    # The 3-month window ending Aug 20, 2026 is strictly May 20 to Aug 20.
    invoices = [
        ['INV-09855', '2025-11-15', 75000, '2025-12-15', '2025-12-18', 33],
        ['INV-09940', '2025-12-10', 82000, '2026-01-09', '2026-01-14', 35],
        ['INV-10101', '2026-02-15', 85000, '2026-03-17', '2026-03-22', 35],
        ['INV-10245', '2026-03-10', 92000, '2026-04-09', '2026-04-17', 38],
        ['INV-10492', '2026-04-05', 88000, '2026-05-05', '2026-05-18', 43],
        ['INV-10611', '2026-05-15', 108740, '2026-06-14', '2026-06-14', 30], # Outside window (if included, average passes)
        ['INV-10884', '2026-05-25', 97325, '2026-06-24', '2026-07-06', 42], # Inside window
        ['INV-11005', '2026-06-05', 141680, '2026-07-05', '2026-07-19', 44], # Inside window
        ['INV-11230', '2026-06-20', 127950, '2026-07-20', '2026-08-07', 48], # Inside window
        ['INV-11350', '2026-07-02', 130000, '2026-08-01', '2026-08-18', 47], # Inside window
        ['INV-11441', '2026-07-15', 150000, '2026-08-14', 'OUTSTANDING', 'N/A'],
        ['INV-11650', '2026-08-01', 150000, '2026-08-31', 'OUTSTANDING', 'N/A'],
        ['INV-11802', '2026-08-10', 172775, '2026-09-09', 'OUTSTANDING', 'N/A']
    ]
    
    for row in invoices:
        ws.append(row)
        
    ws.append(['', '', '', '', '', ''])
    ws.append(["Note: 'Days to Pay' represents the number of days between Invoice Date and Date Paid.", '', '', '', '', ''])
    
    # Sheet 2: Account Summary
    ws_sum = wb.create_sheet(title="Account Summary")
    ws_sum.append(['Hartsfield Polymer Technologies, LLC - Account Summary', ''])
    ws_sum.append(['', ''])
    ws_sum.append(['Credit Limit', 400000])
    ws_sum.append(['Current Balance', 472775]) # Sum of 11441, 11650, 11802 = 150000+150000+172775
    ws_sum.append(['Available Credit', -72775])
    ws_sum.append(['', ''])
    
    ws_sum.append(['Aging Summary', 'Amount (USD)'])
    ws_sum.append(['Current (Not Yet Due)', 322775]) # 11650, 11802
    ws_sum.append(['1-30 Days Past Due', 150000]) # 11441
    ws_sum.append(['31-60 Days Past Due', 0])
    ws_sum.append(['61-90 Days Past Due', 0])
    ws_sum.append(['91+ Days Past Due', 0])
    ws_sum.append(['Total Outstanding', 472775])
    
    # Format
    for ws_cur in [ws, ws_sum]:
        for col in ['A', 'B', 'C', 'D', 'E', 'F']:
            ws_cur.column_dimensions[col].width = 18
            
    wb.save('Crestline_AR_Ledger_Hartsfield_History.xlsx')

if __name__ == '__main__':
    create_financials()
    create_ar_ledger()
    print("Excel files generated successfully.")
