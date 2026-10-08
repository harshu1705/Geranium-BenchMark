import sys
import re
from docx import Document
import openpyxl

def verify_golden():
    doc = Document('Trade_Credit_Decision_Memo.docx')
    text = ' '.join([p.text for p in doc.paragraphs])
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text += ' ' + cell.text
                
    checks = {
        'August 20, 2026': False,
        'Decline': False,
        '1.77': False,
        '45.25': False,
        '$805,450': False, # EBIT check
        'INV-10611': False, # Check that the edge case is explicitly excluded
        'fails two or more': False # check deterministic section 4 text
    }

    for key in checks.keys():
        if key in text:
            checks[key] = True
        else:
            print(f"FAILED: Missing '{key}' in Golden Solution")

    print("\n--- GOLDEN MATRIX ---")
    for k, v in checks.items():
        print(f"{k}: {'PASS' if v else 'FAIL'}")

def verify_excel():
    wb_ar = openpyxl.load_workbook('Crestline_AR_Ledger_Hartsfield_History.xlsx', data_only=True)
    ws_ar = wb_ar.active
    
    # check that INV-10611 is present and has date 2026-05-15
    found = False
    for row in ws_ar.iter_rows(values_only=True):
        if row and row[0] == 'INV-10611' and row[1] == '2026-05-15':
            found = True
            
    print(f"INV-10611 (Out of Window Trap) present in Excel: {'PASS' if found else 'FAIL'}")
    
    wb_fin = openpyxl.load_workbook('Hartsfield_Polymer_Unaudited_Financials.xlsx', data_only=True)
    ws_is = wb_fin["Income Statement"]
    ebit_found = False
    for row in ws_is.iter_rows(values_only=True):
        if row and row[0] == 'Earnings Before Interest and Taxes (EBIT)' and row[1] == 805450:
            ebit_found = True
    print(f"EBIT = $805,450 present in Excel: {'PASS' if ebit_found else 'FAIL'}")

if __name__ == '__main__':
    verify_golden()
    verify_excel()
