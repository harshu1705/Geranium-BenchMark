import docx
import shutil

SRC = "golden_extract/Trade_Credit_Decision_Memo.docx"
DST = "golden_fix_final7/Trade_Credit_Decision_Memo.docx"
shutil.copy(SRC, DST)

d = docx.Document(DST)
paras = d.paragraphs

replacements = {
    2: ("To: Corporate Credit Committee, Crestline Chemical Supply, Inc.",
        "To: Corporate Credit Department, Crestline Chemical Supply, Inc."),
    3: ("From: Credit Analysis Team",
        "From: Credit Analyst"),
}

for idx, (old, new) in replacements.items():
    p = paras[idx]
    assert len(p.runs) == 1, (idx, [r.text for r in p.runs])
    assert p.runs[0].text == old, (idx, p.runs[0].text)
    p.runs[0].text = new

# Para 7: second run (non-bold) contains the Section 4(b) reference
p7 = paras[7]
assert len(p7.runs) == 2
old7 = p7.runs[1].text
assert "Section 4(b)" in old7
new7 = old7.replace(
    "Under Section 4(b) of the credit policy, because",
    'Under Section 4\'s "b) Decline the Request (2+ Failures or Tax Lien)" provision, because'
)
assert new7 != old7
p7.runs[1].text = new7

# Para 22: tax lien date format
p22 = paras[22]
assert len(p22.runs) == 1
old22 = p22.runs[0].text
assert "July 12, 2026" in old22
new22 = old22.replace("filed on July 12, 2026", "filed on 7/12/2026")
assert new22 != old22
p22.runs[0].text = new22

# Para 25: conclusion Section 4(b) reference
p25 = paras[25]
assert len(p25.runs) == 1
old25 = p25.runs[0].text
assert "Section 4(b)" in old25
new25 = old25.replace(
    "Under Section 4(b) of the credit policy, because",
    'Under Section 4\'s "b) Decline the Request (2+ Failures or Tax Lien)" provision, because'
)
assert new25 != old25
p25.runs[0].text = new25

d.save(DST)
print("Saved:", DST)

# Verify
d2 = docx.Document(DST)
for i in [2, 3, 7, 22, 25]:
    print(i, repr(d2.paragraphs[i].text))
