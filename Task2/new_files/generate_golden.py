import os
from docx import Document
from docx.shared import Pt, Inches

def create_golden():
    doc = Document()
    
    doc.add_heading('Trade Credit Decision Memo', 0)
    doc.add_paragraph('Date: August 20, 2026')
    doc.add_paragraph('To: Corporate Credit Committee, Crestline Chemical Supply, Inc.')
    doc.add_paragraph('From: Credit Analysis Team')
    doc.add_paragraph('Subject: Hartsfield Polymer Technologies, LLC - Credit Limit Request')
    
    doc.add_heading('Executive Summary', level=1)
    doc.add_paragraph(
        'Hartsfield Polymer Technologies, LLC has requested an increase to their unsecured trade credit limit from $400,000 (Net 30) to $1,500,000 (Net 60). '
        'Based on a comprehensive and strict review of their financial statements, payment history, and external credit reports, the applicant fails all three mandatory criteria under Section 3 of the Crestline Corporate Credit Policy.'
    )
    p = doc.add_paragraph()
    p.add_run('Recommendation: Decline. ').bold = True
    p.add_run(
        'Under the revised Section 4 of the credit policy, because Hartsfield Polymer fails two or more mandatory criteria (and possesses an active tax lien), '
        'the analyst must strictly decline the limit increase and maintain the customer on their existing limit. '
        'Conditional Approval via CRM is explicitly prohibited in this scenario.'
    )
    
    doc.add_heading('Independent Financial Health Assessment', level=1)
    doc.add_paragraph(
        'A review of the 2025 unaudited financials reveals a leveraged balance sheet with weak liquidity. '
        'The Current Ratio is 1.29 ($4,115,200 / $3,184,500), but the Quick Ratio is lower at 0.63 (($485,200 + $1,515,000) / $3,184,500), indicating heavy reliance on inventory. '
        'Leverage is extremely high, with a Debt-to-Equity ratio of 2.47 ($7,245,600 / $2,936,900). '
        'Profitability margins are thin; the 2025 Net Profit Margin was 4.76% ($496,950 / $10,435,200). '
        'The H1 2026 interim statements show continued thin margins and do not support a $1.5M unsecured exposure.'
    )
    
    doc.add_heading('Quantitative Financial Analysis (Z-Score)', level=1)
    doc.add_paragraph(
        'Per Section 3 of the policy, unsecured requests over $1,000,000 require an Altman Z-Score of 1.80 or higher. Based on the 2025 unaudited financials:'
    )
    
    doc.add_paragraph(
        'Total Assets = $10,182,500\n'
        'Total Liabilities = $7,245,600\n'
        'Working Capital = Current Assets ($4,115,200) - Current Liabilities ($3,184,500) = $930,700'
    )
    
    doc.add_paragraph(
        'A = Working Capital / Total Assets = $930,700 / $10,182,500 = 0.0914\n'
        'B = Retained Earnings / Total Assets = $936,900 / $10,182,500 = 0.0920\n'
        'C = EBIT / Total Assets = $805,450 / $10,182,500 = 0.0791\n'
        'D = Book Value of Equity / Total Liabilities = $2,936,900 / $7,245,600 = 0.4053\n'
        'E = Sales / Total Assets = $10,435,200 / $10,182,500 = 1.0248'
    )
    
    doc.add_paragraph(
        'Calculated Z-Score = 1.2(0.0914) + 1.4(0.0920) + 3.3(0.0791) + 0.6(0.4053) + 1.0(1.0248) = 0.1097 + 0.1288 + 0.2610 + 0.2432 + 1.0248 = 1.7675.'
    )
    
    doc.add_paragraph(
        'The calculated Z-score of 1.77 is below the 1.80 policy minimum, triggering a mandatory policy failure.'
    )
    
    doc.add_heading('Internal Payment History', level=1)
    doc.add_paragraph(
        "The policy requires that average 'Days to Pay' over the most recent 3-month period not exceed 45 days. "
        "The decision date is August 20, 2026, meaning the strict 3-month window covers invoices dated from May 20, 2026, through August 20, 2026. "
        "Invoice INV-10611 (dated May 15) falls outside this window and must be excluded. "
        "Reviewing the Crestline AR Ledger for Hartsfield within the strict window, we observe the following paid invoices:"
    )
    
    doc.add_paragraph(
        '- INV-10884: Paid in 42 days.\n'
        '- INV-11005: Paid in 44 days.\n'
        '- INV-11230: Paid in 48 days.\n'
        '- INV-11350: Paid in 47 days.'
    )
    
    doc.add_paragraph(
        'Average Days to Pay = (42 + 44 + 48 + 47) / 4 = 45.25 days.'
    )
    
    doc.add_paragraph(
        "Hartsfield's average days to pay over the exact 3-month window is 45.25 days, which strictly exceeds the 45-day maximum threshold. "
        "This constitutes a second policy failure."
    )
    
    doc.add_heading('Public Records and External Credit Assessment', level=1)
    doc.add_paragraph(
        'The D&B Credit Report shows an active State Tax Lien filed on July 12, 2026, for $45,000.'
    )
    
    doc.add_paragraph(
        'Section 3 of the credit policy strictly prohibits active or unresolved tax liens exceeding $10,000. '
        'Because this lien remains active on public record, it triggers a third mandatory policy failure.'
    )
    
    doc.add_heading('Conclusion', level=1)
    doc.add_paragraph(
        'Hartsfield Polymer Technologies strictly failed all three mandatory criteria for a $1,000,000+ unsecured line. '
        'They missed the Z-score minimum (1.77 vs 1.80), exceeded the allowable internal days-to-pay metric within the 3-month window (45.25 vs 45), and carry an unresolved $45,000 tax lien. '
        'Under Section 4(b) of the credit policy, because the applicant failed two or more criteria (and carries an active tax lien), the limit increase must be outright Declined, with no option for conditional approval.'
    )
    
    doc.save('Trade_Credit_Decision_Memo.docx')

if __name__ == '__main__':
    create_golden()
    print("Golden solution generated successfully.")
