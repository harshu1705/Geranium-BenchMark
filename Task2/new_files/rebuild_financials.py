import openpyxl

wb = openpyxl.load_workbook('Hartsfield_Polymer_Unaudited_Financials.xlsx')

CUR = '#,##0'

def set_cur(ws, coord):
    ws[coord].number_format = CUR

# ---- Income Statement (FY2025) ----
ws = wb['Income Statement']
ws['B7'] = '=B5-B6'                 # Gross Profit
ws['B10'] = '=B7-B8-B9'             # EBIT
ws['B12'] = '=B10-B11'              # EBT
ws['B14'] = '=B12-B13'              # Net Income
for r in [5, 6, 7, 8, 9, 10, 11, 12, 13, 14]:
    set_cur(ws, f'B{r}')

# ---- Balance Sheet (FY2025 col B, FY2024 col C) ----
ws = wb['Balance Sheet']
for col in ['B', 'C']:
    ws[f'{col}10'] = f'=SUM({col}6:{col}9)'      # Total Current Assets
    ws[f'{col}13'] = f'={col}10+{col}11+{col}12'  # Total Assets
    ws[f'{col}19'] = f'=SUM({col}16:{col}18)'     # Total Current Liabilities
    ws[f'{col}21'] = f'={col}19+{col}20'          # Total Liabilities
    ws[f'{col}26'] = f'={col}24+{col}25'          # Total Equity (Book Value)
    ws[f'{col}27'] = f'={col}21+{col}26'          # Total Liabilities & Equity
    for r in range(6, 28):
        if ws[f'{col}{r}'].value is not None:
            set_cur(ws, f'{col}{r}')

# ---- H1 2026 Interim Income Statement ----
ws = wb['H1 2026 Interim IS']
ws['B7'] = '=B5-B6'
ws['B10'] = '=B7-B8-B9'
ws['B12'] = '=B10-B11'
ws['B14'] = '=B12-B13'
for r in [5, 6, 7, 8, 9, 10, 11, 12, 13, 14]:
    set_cur(ws, f'B{r}')

wb.save('Hartsfield_Polymer_Unaudited_Financials.xlsx')
print('saved')
