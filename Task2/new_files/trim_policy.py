import docx
d = docx.Document('Crestline_Corporate_Credit_Policy.docx')

p = d.paragraphs[10]
old = p.text
assert old.startswith('2. Internal Payment History'), old

new_text = (
    "2. Internal Payment History: For existing accounts requesting limit increases, payment performance "
    "must be satisfactory. Their average 'Days to Pay' over the most recent 3-month period on the Crestline "
    "Accounts Receivable (AR) Ledger must not exceed 45 days, considering only invoices falling within that "
    "3-month window immediately preceding the decision date. Accounts exhibiting a trend of slowing payments "
    "beyond this 45-day threshold are strictly ineligible for limit increases on unsecured terms."
)

for run in list(p.runs):
    run.text = ''
if p.runs:
    p.runs[0].text = new_text
else:
    p.add_run(new_text)

d.save('Crestline_Corporate_Credit_Policy.docx')
print('saved')
