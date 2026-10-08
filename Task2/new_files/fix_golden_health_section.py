from docx import Document

d = Document('Trade_Credit_Decision_Memo.docx')

target = None
for p in d.paragraphs:
    if p.text.startswith('A review of the 2025 unaudited financials'):
        target = p
        break
assert target is not None

new_text = (
    "A review of the 2025 unaudited financials reveals a leveraged balance sheet with weak liquidity. "
    "The Current Ratio is 1.29 ($4,115,200 / $3,184,500), but the Quick Ratio is lower at 0.63 "
    "(($485,200 + $1,515,000) / $3,184,500), indicating heavy reliance on inventory. Leverage is extremely "
    "high, with a Debt-to-Equity ratio of 2.47 ($7,245,600 / $2,936,900). Profitability margins are thin; "
    "the 2025 Net Profit Margin was 4.76% ($496,950 / $10,435,200). "
    "The H1 2026 interim statements confirm this deterioration is accelerating rather than stabilizing: "
    "EBIT margin fell to 2.48% ($145,600 / $5,865,200) from 7.72% for full-year 2025 ($805,450 / $10,435,200), "
    "a decline of roughly two-thirds. Interest coverage (EBIT / Interest Expense) fell to 1.08x "
    "($145,600 / $135,000) for the six months ended June 30, 2026, versus 7.29x for full-year 2025 "
    "($805,450 / $110,500), meaning operating earnings now barely cover interest expense. This interim "
    "trend is the strongest evidence in the file that Hartsfield's financial health does not support a "
    "$1.5M unsecured exposure, independent of the Section 3 Z-Score, payment history, and public records gates "
    "addressed below."
)

for r in target.runs:
    r.text = ''
if target.runs:
    target.runs[0].text = new_text
else:
    target.add_run(new_text)

d.save('Trade_Credit_Decision_Memo.docx')
print('saved')
