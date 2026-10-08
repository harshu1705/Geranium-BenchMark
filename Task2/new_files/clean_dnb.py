# -*- coding: utf-8 -*-
import docx
from docx import Document

d = Document('Dun_and_Bradstreet_Credit_Report.docx')

def set_text(p, text):
    for run in list(p.runs):
        run.text = ''
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run(text)

def delete_paragraph(p):
    p._element.getparent().remove(p._element)

for p in list(d.paragraphs):
    t = p.text.strip()

    if t.startswith('D&B Paydex Score: 68'):
        set_text(p, "A Paydex of 68 indicates that payments to suppliers are typically made 15 to 22 days "
                     "beyond standard terms. The Paydex score is a unique, dollar-weighted indicator of a "
                     "business's past payment performance based on trade experiences reported to Dun & Bradstreet "
                     "by various vendors. A score of 80 denotes that payments are made exactly within terms, "
                     "while scores below 70 generally flag an elevated risk of late payments.")

    elif t.startswith('Financial Stress Score: 1450'):
        set_text(p, "This indicates a higher than average risk of experiencing severe financial stress over the "
                     "next 12 months compared to similar businesses in the manufacturing sector. The Financial "
                     "Stress Score utilizes statistical modeling to predict the likelihood of business failure, "
                     "combining demographic data, financial ratios, and payment history.")

    elif t.startswith('Based on 14 trade experiences'):
        set_text(p, "Based on 14 trade experiences reported in the last 12 months, the company's payment "
                     "performance breaks down as follows:")

    elif t.startswith('1. Raw Materials Supplier'):
        delete_paragraph(p)

    elif t in ('Bankruptcies: 0 recorded.', 'Judgments: 0 recorded.', 'Liens: 1 active record found.') \
            or t.startswith('- Filed Date:'):
        delete_paragraph(p)

d.save('Dun_and_Bradstreet_Credit_Report.docx')
print('cleaned')
