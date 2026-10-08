from docx import Document

d = Document('Crestline_Corporate_Credit_Policy.docx')
target = d.paragraphs[10]
assert 'Days to Pay' in target.text

new_text = (
    "2. Internal Payment History: For existing accounts requesting limit increases, payment performance "
    "must be satisfactory. Their average 'Days to Pay' over the most recent 3-month period on the Crestline "
    "Accounts Receivable (AR) Ledger must not exceed 45 days. This 3-month window is measured by each "
    "invoice's Invoice Date, not its Date Paid; only invoices dated within the 3 months immediately preceding "
    "the decision date are included in the average. Accounts exhibiting a trend of slowing payments beyond "
    "this 45-day threshold are strictly ineligible for limit increases on unsecured terms."
)

for r in target.runs:
    r.text = ''
if target.runs:
    target.runs[0].text = new_text
else:
    target.add_run(new_text)

d.save('Crestline_Corporate_Credit_Policy.docx')
print('saved')
