import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
import openpyxl
import os

BASE = r'c:\Users\user\OneDrive\Desktop\geranium project\Task4'
INDIR = os.path.join(BASE, 'inputs')
os.makedirs(INDIR, exist_ok=True)

# ---------------------------------------------------------------------------
# 1. Ellison_Financial_Planning_Guidelines.docx
# ---------------------------------------------------------------------------
d = Document()
d.add_heading('MERIDIAN WEALTH PARTNERS — RETIREMENT INCOME PLANNING METHODOLOGY', level=0)
d.add_paragraph('Internal Planning Standards | Effective January 2026')

d.add_heading('1. Purpose', level=1)
d.add_paragraph(
    "This document sets the methodology advisors must follow when preparing a Retirement Income Plan "
    "for a client household. Every assumption below is a firm-level standard and must be applied exactly "
    "as written; advisors may not substitute personal judgment for any numeric threshold in this document."
)

d.add_heading('2. Safe Withdrawal Rate Standard', level=1)
d.add_paragraph(
    "The firm's standard sustainable withdrawal rate is 4.0% of the client's total investable portfolio "
    "balance, excluding any cash held as an emergency reserve. If a client's calculated income gap requires "
    "a withdrawal rate that exceeds 4.0% by more than 1.0 percentage point, the plan must flag this "
    "explicitly as an elevated-risk withdrawal requiring documented client acknowledgment of increased "
    "portfolio depletion risk before implementation."
)

d.add_heading('3. Asset Allocation and Rebalancing', level=1)
d.add_paragraph(
    "Target allocations by risk tolerance: Conservative 40% equity / 60% bond; Moderate 60% equity / 40% "
    "bond; Growth 80% equity / 20% bond. A client's current blended allocation is calculated by weighting "
    "each account's stated equity/bond split by its account balance across the full investable portfolio. "
    "If the blended equity percentage drifts more than 5 percentage points above or below the target in "
    "either direction, the account is outside its rebalancing band and must be rebalanced back to the "
    "target percentage (not merely to the edge of the band)."
)

d.add_heading('4. Social Security Claiming Analysis', level=1)
d.add_paragraph(
    "For each spouse, calculate the breakeven age between claiming at 62 and claiming at 70 using simple, "
    "undiscounted cumulative benefits (do not apply a discount rate or present-value adjustment). The "
    "breakeven age is the age at which cumulative benefits under delayed claiming first equal cumulative "
    "benefits under early claiming. Compare the breakeven age to the client's planning life expectancy: 92 "
    "for male clients, 94 for female clients, reflecting the firm's standard longevity planning "
    "convention. If the client's planning life expectancy exceeds the breakeven age, recommend delaying to "
    "age 70; otherwise recommend claiming at 62."
)

d.add_heading('5. Retirement Income Gap Analysis', level=1)
d.add_paragraph(
    "The annual income gap is the client's stated desired annual retirement spending minus all guaranteed "
    "income currently being received or scheduled to begin immediately (pensions already in pay status). "
    "Do not include Social Security benefits in guaranteed income for this calculation if the recommended "
    "claiming strategy under Section 4 defers the start of benefits — the gap must reflect the bridge "
    "period before Social Security begins. Compare the gap to the sustainable withdrawal amount calculated "
    "under Section 2 to determine any shortfall or surplus."
)

d.add_heading('6. Withdrawal Sequencing', level=1)
d.add_paragraph(
    "When funding the income gap from the portfolio, withdraw first from taxable brokerage accounts, then "
    "from tax-deferred accounts, and preserve Roth accounts for last. For this engagement, treat "
    "withdrawals from a taxable brokerage account as return of basis (not additional ordinary taxable "
    "income) unless the taxable account balance is insufficient to cover the full withdrawal need for the "
    "year, in which case only the remaining, unfunded portion is treated as an additional tax-deferred "
    "withdrawal and taxed as ordinary income."
)

d.add_heading('7. Roth Conversion Opportunity', level=1)
d.add_paragraph(
    "During any year in which a client's Social Security benefits have not yet begun (a bridge year), "
    "calculate the room available to convert Traditional IRA or 401(k) assets to Roth by filling remaining "
    "space up to the top of the 12% federal ordinary income tax bracket. Use the standard deduction and "
    "bracket thresholds provided in Federal_Tax_Bracket_Assumptions_2026.xlsx for this engagement; do not "
    "look up or apply any other year's tax law. Include all ordinary taxable income for the year (pension "
    "income, taxable interest and non-qualified dividends, and any tax-deferred withdrawal required under "
    "Section 6) before calculating remaining bracket room. The conversion amount recommended may not exceed "
    "the balance of the account being converted."
)

d.add_heading('8. Deliverable', level=1)
d.add_paragraph(
    "Each client engagement produces a single Retirement Income Plan document covering: the client's "
    "current financial position, the Social Security claiming recommendation, the income gap and "
    "withdrawal analysis, the Roth conversion opportunity (if any), the asset allocation review, and a "
    "final set of recommendations with the assumptions used clearly disclosed."
)

d.save(os.path.join(INDIR, 'Ellison_Financial_Planning_Guidelines.docx'))
print('Saved guidelines doc')

# ---------------------------------------------------------------------------
# 2. Ellison_Account_Statements_Q4_2025.xlsx
# ---------------------------------------------------------------------------
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Accounts'
ws.append(['Account Owner', 'Account Name', 'Tax Treatment', 'Balance ($)', 'Equity %', 'Bond %'])
rows = [
    ('David', "David 401(k)", 'Tax-Deferred', 612400, 0.70, 0.30),
    ('David', 'David Traditional IRA', 'Tax-Deferred', 184900, 0.65, 0.35),
    ('David', 'David Roth IRA', 'Tax-Free', 76300, 0.80, 0.20),
    ('Karen', 'Karen 401(k)', 'Tax-Deferred', 298650, 0.60, 0.40),
    ('Karen', 'Karen Roth IRA', 'Tax-Free', 52100, 0.75, 0.25),
    ('Joint', 'Joint Taxable Brokerage', 'Taxable', 341750, 0.65, 0.35),
    ('Joint', 'Joint Savings / Money Market (Emergency Reserve)', 'Cash Reserve', 45000, 0.00, 1.00),
]
for r in rows:
    ws.append(r)
ws.append(['', 'Note: Emergency Reserve is held outside the investable portfolio per Guidelines Section 2 and is excluded from the withdrawal-rate and allocation calculations.', '', '', '', ''])
wb.save(os.path.join(INDIR, 'Ellison_Account_Statements_Q4_2025.xlsx'))
print('Saved account statements')

# ---------------------------------------------------------------------------
# 3. Ellison_Social_Security_Benefit_Estimates.xlsx
# ---------------------------------------------------------------------------
wb2 = openpyxl.Workbook()
ws2 = wb2.active
ws2.title = 'SSA Estimates'
ws2.append(['Spouse', 'Date of Birth', 'Claiming Age', 'Estimated Monthly Benefit ($)'])
data = [
    ('David', '1963-11-04', 62, 2340),
    ('David', '1963-11-04', 67, 3380),
    ('David', '1963-11-04', 70, 4230),
    ('Karen', '1965-03-19', 62, 1290),
    ('Karen', '1965-03-19', 67, 1860),
    ('Karen', '1965-03-19', 70, 2330),
]
for r in data:
    ws2.append(r)
wb2.save(os.path.join(INDIR, 'Ellison_Social_Security_Benefit_Estimates.xlsx'))
print('Saved SSA estimates')

# ---------------------------------------------------------------------------
# 4. Ellison_Client_Goals_and_Income_Summary.docx
# ---------------------------------------------------------------------------
d2 = Document()
d2.add_heading('Client Goals and Income Summary — David & Karen Ellison', level=0)
d2.add_paragraph('Prepared for: Meridian Wealth Partners | Engagement Date: January 2026')

d2.add_heading('Household', level=1)
d2.add_paragraph('David Ellison — DOB 11/04/1963 — currently age 62 — retired as of December 31, 2025.')
d2.add_paragraph('Karen Ellison — DOB 03/19/1965 — currently age 60 — retired as of December 31, 2025.')
d2.add_paragraph('Filing status: Married Filing Jointly.')

d2.add_heading('Retirement Goals', level=1)
d2.add_paragraph('Desired annual retirement spending (after-tax): $110,000.')
d2.add_paragraph('Risk tolerance (per completed questionnaire): Moderate.')
d2.add_paragraph('Legacy goal: preserve at least $150,000 for children; not a hard constraint on this analysis but noted for future reviews.')

d2.add_heading('Guaranteed Income Sources', level=1)
d2.add_paragraph('David has a small defined-benefit pension from a prior employer, already in pay status: $1,100 per month ($13,200 per year), fully taxable as ordinary income. Karen has no pension.')
d2.add_paragraph('Neither spouse has yet filed for Social Security. Both are considering when to claim.')

d2.add_heading('Other Taxable Income', level=1)
d2.add_paragraph('The joint taxable brokerage account is expected to generate approximately $6,200 in taxable interest and non-qualified dividend income in the current year, independent of any withdrawals of principal.')

d2.save(os.path.join(INDIR, 'Ellison_Client_Goals_and_Income_Summary.docx'))
print('Saved client goals doc')

# ---------------------------------------------------------------------------
# 5. Federal_Tax_Bracket_Assumptions_2026.xlsx
# ---------------------------------------------------------------------------
wb3 = openpyxl.Workbook()
ws3 = wb3.active
ws3.title = 'MFJ Brackets'
ws3.append(['Filing Status', 'Standard Deduction ($)'])
ws3.append(['Married Filing Jointly', 30000])
ws3.append([])
ws3.append(['Ordinary Bracket', 'Taxable Income From ($)', 'Taxable Income To ($)', 'Rate'])
brackets = [
    ('10%', 0, 23850, 0.10),
    ('12%', 23850, 96950, 0.12),
    ('22%', 96950, 206700, 0.22),
]
for r in brackets:
    ws3.append(r)
ws3.append([])
ws3.append(['Note: These bracket thresholds and the standard deduction are firm-provided assumptions for this engagement only. Do not substitute actual IRS figures for any other year.'])
wb3.save(os.path.join(INDIR, 'Federal_Tax_Bracket_Assumptions_2026.xlsx'))
print('Saved tax bracket assumptions')

print('\nAll 5 input files created in', INDIR)
