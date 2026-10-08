import os
import openpyxl
from docx import Document

def dump_docx(filepath):
    print(f"\n{'='*50}\nFILE: {filepath}\n{'='*50}")
    try:
        doc = Document(filepath)
        for p in doc.paragraphs:
            print(p.text)
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    row_text.append(cell.text.replace('\n', ' '))
                print(' | '.join(row_text))
    except Exception as e:
        print(f"Error reading {filepath}: {e}")

def dump_xlsx(filepath):
    print(f"\n{'='*50}\nFILE: {filepath}\n{'='*50}")
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        for ws in wb.worksheets:
            print(f"\n--- SHEET: {ws.title} ---")
            for row in ws.iter_rows(values_only=True):
                print([str(cell) if cell is not None else '' for cell in row])
    except Exception as e:
        print(f"Error reading {filepath}: {e}")

input_dir = r"C:\Users\user\OneDrive\Desktop\geranium project\NovaMart_extracted\Inputs"
for file in os.listdir(input_dir):
    filepath = os.path.join(input_dir, file)
    if file.endswith('.docx'):
        dump_docx(filepath)
    elif file.endswith('.xlsx'):
        dump_xlsx(filepath)

golden_dir = r"C:\Users\user\OneDrive\Desktop\geranium project\NovaMart_extracted\Golden"
for file in os.listdir(golden_dir):
    filepath = os.path.join(golden_dir, file)
    if file.endswith('.docx'):
        dump_docx(filepath)
