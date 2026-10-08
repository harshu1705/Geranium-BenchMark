import zipfile
from docx import Document
import openpyxl

print("=" * 70)
print("1. METADATA CHECK (no generator leakage)")
print("=" * 70)
docx_files = [
    'Inputs/Broker_Email_Correspondence.docx',
    'Inputs/NovaMart_Cyber_Application.docx',
    'Inputs/Ransomware_Supplemental_App.docx',
    'Inputs/Underwriting_Guidelines_v2.docx',
    'Golden/Underwriting_Memo_NovaMart.docx',
]
xlsx_files = [
    'Inputs/Financial_Statements_Summary.xlsx',
    'Inputs/Loss_Run_History_2021_2026.xlsx',
    'Inputs/Third_Party_Security_Scan.xlsx',
]
for f in docx_files:
    d = Document(f)
    cp = d.core_properties
    bad = 'python-docx' in (cp.author or '') or 'python-docx' in (cp.comments or '')
    print(f"  {'FAIL' if bad else 'OK  '} {f} | author={cp.author!r} lmb={cp.last_modified_by!r} comments={cp.comments!r}")

for f in xlsx_files:
    with zipfile.ZipFile(f) as z:
        app = z.read('docProps/app.xml').decode('utf-8')
    bad = 'openpyxl' in app.lower()
    print(f"  {'FAIL' if bad else 'OK  '} {f} | app.xml has openpyxl marker: {bad}")

print()
print("=" * 70)
print("2. WORD / CELL COUNT CHECK (thin file fix)")
print("=" * 70)
d = Document('Inputs/NovaMart_Cyber_Application.docx')
words = sum(len(p.text.split()) for p in d.paragraphs)
print(f"  NovaMart_Cyber_Application.docx: {words} words (was ~205)")

wb = openpyxl.load_workbook('Inputs/Financial_Statements_Summary.xlsx', data_only=True)
ws = wb.active
cells = sum(1 for row in ws.iter_rows() for c in row if c.value not in (None, '-'))
print(f"  Financial_Statements_Summary.xlsx: {cells} non-dash populated cells (was ~80)")

wb2 = openpyxl.load_workbook('Inputs/Loss_Run_History_2021_2026.xlsx', data_only=True)
ws2 = wb2.active
cells2 = sum(1 for row in ws2.iter_rows() for c in row if c.value not in (None, '-'))
print(f"  Loss_Run_History_2021_2026.xlsx: {cells2} non-dash populated cells (was ~95)")

print()
print("=" * 70)
print("3. FORMULA CHECK")
print("=" * 70)
wbf = openpyxl.load_workbook('Inputs/Financial_Statements_Summary.xlsx', data_only=False)
wsf = wbf.active
formula_cells = [c.coordinate for row in wsf.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith('=')]
print(f"  Financial_Statements_Summary.xlsx formula cells: {len(formula_cells)} -> {formula_cells[:8]}...")

wbl = openpyxl.load_workbook('Inputs/Loss_Run_History_2021_2026.xlsx', data_only=False)
wsl = wbl.active
formula_cells2 = [c.coordinate for row in wsl.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith('=')]
print(f"  Loss_Run_History_2021_2026.xlsx formula cells: {formula_cells2}")

print()
print("=" * 70)
print("4. CROSS-DOCUMENT CONSISTENCY CHECK (golden vs inputs)")
print("=" * 70)
gold = Document('Golden/Underwriting_Memo_NovaMart.docx')
gtext = '\n'.join(p.text for p in gold.paragraphs)
for t in gold.tables:
    for r in t.rows:
        for c in r.cells:
            gtext += '\n' + c.text

checks = {
    '$502,400,000 (TTM revenue)': '$502,400,000' in gtext,
    '$25,120 (base premium)': '$25,120' in gtext,
    '$650,000 (cyber-only loss total)': '$650,000' in gtext,
    '$37,680 (final premium)': '$37,680' in gtext,
    '1.20x multiplier': '1.20x' in gtext,
    '1.25x multiplier': '1.25x' in gtext,
    '$10,000,000 cap': '$10,000,000' in gtext,
    '$15,000,000 requested': '$15,000,000' in gtext,
    'CVE-2023-4528 (open)': 'CVE-2023-4528' in gtext,
    'CVE-2024-3094 (remediated)': 'CVE-2024-3094' in gtext,
    'CVE-2022-40684 (remediated, NEW)': 'CVE-2022-40684' in gtext,
    'Senior Underwriter authority': 'Senior Underwriter' in gtext,
    'DECLINE recommendation': 'DECLINE' in gtext,
    'No "multiple simultaneous decline" framing': 'multiple simultaneous decline' not in gtext,
    'No stale $657,500 figure': '657,500' not in gtext,
    'No stale $722,000 figure': '722,000' not in gtext,
    'No "uncompetitive" framing': 'uncompetitive' not in gtext,
    'Broker target $45,000 referenced': '$45,000' in gtext,
}
for k, v in checks.items():
    print(f"  {'PASS' if v else 'FAIL'}  {k}")

print()
print("=" * 70)
print("5. SCAN FILE - confirm CVE-2022-40684 really is Critical/9.6/VPN")
print("=" * 70)
wbs = openpyxl.load_workbook('Inputs/Third_Party_Security_Scan.xlsx', data_only=True)
wss = wbs.active
for row in wss.iter_rows(min_row=2, values_only=True):
    if row[4] == 'CVE-2022-40684':
        print(f"  Found: {row}")

print()
print("=" * 70)
print("6. LOSS RUN - confirm 3-year window math (cyber-only vs all-lines)")
print("=" * 70)
wbl2 = openpyxl.load_workbook('Inputs/Loss_Run_History_2021_2026.xlsx', data_only=True)
wsl2 = wbl2.active
from datetime import date
window_start = date(2023, 10, 1)
cyber_total = 0
all_total = 0
for row in wsl2.iter_rows(min_row=2, max_row=11, values_only=True):
    claim_id, dol, lob, desc, status, legal, forensic, ransom, bi, total_paid, reserve, claimant, adjuster = row
    d_ = date.fromisoformat(dol)
    if d_ >= window_start:
        all_total += total_paid
        if lob == 'Cyber':
            cyber_total += total_paid
print(f"  All-lines in-window total: ${all_total:,}")
print(f"  Cyber-only in-window total: ${cyber_total:,}")
