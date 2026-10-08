from docx import Document

path = 'Inputs/NovaMart_Cyber_Application.docx'
d = Document(path)


def insert_after(anchor_para, doc, heading=None, level=2, body_paras=None):
    """Insert a heading + body paragraphs immediately after anchor_para."""
    last = anchor_para
    if heading:
        h = doc.add_paragraph(heading, style=f'Heading {level}')
        last._p.addnext(h._p)
        last = h
    for text in (body_paras or []):
        b = doc.add_paragraph(text)
        last._p.addnext(b._p)
        last = b
    return last


paras = d.paragraphs

# 1. After basic info block (index 1) - add corporate structure details
anchor = paras[1]
insert_after(anchor, d, heading='Corporate & Operational Profile', level=2, body_paras=[
    'Years in Business: 22 years (founded 2004). Legal Structure: C-Corporation, privately held, no publicly traded securities.',
    'Number of Employees: Approximately 6,100 full-time and part-time employees across all locations, including approximately 140 corporate/IT staff at the Chicago headquarters.',
    'Store Footprint: 45 locations across Illinois, Indiana, Wisconsin, and Michigan, plus one central distribution warehouse in Joliet, IL and a secondary overflow warehouse acquired with ValueStore in Rockford, IL.',
    'Parent/Subsidiary Relationships: NovaMart Inc. is the sole named insured. ValueStore Holdings LLC, acquired in Q1 2025, now operates as a wholly owned, non-incorporated division of NovaMart Inc. and is not a separate legal entity for insurance purposes.',
])

paras = d.paragraphs
biz_overview_idx = None
for i, p in enumerate(paras):
    if p.text.startswith('NovaMart Inc. is a regional big-box retail chain'):
        biz_overview_idx = i
        break
anchor = paras[biz_overview_idx]
insert_after(anchor, d, heading='IT Infrastructure & Network Environment', level=2, body_paras=[
    'Describe your primary network architecture, including cloud usage, point-of-sale systems, and remote access.',
    'Applicant Response: Our corporate environment runs on a hybrid model: core ERP (SAP) and financial systems are hosted on-premise at our Chicago data center, while e-commerce, loyalty program, and analytics workloads run in AWS (us-east-1 and us-east-2). Point-of-sale terminals at all 45 stores run a PCI-validated P2PE (point-to-point encryption) solution from a third-party vendor, so cardholder data is encrypted at the point of swipe/tap and never traverses our internal network in cleartext. Store networks are segmented from corporate IT via next-generation firewalls, and the ValueStore legacy environment (acquired 2025) remains on a quarantined VLAN pending full migration into our central Active Directory, expected to complete in Q3 2027.',
    'How many servers, workstations, and point-of-sale terminals are in scope for this policy?',
    'Applicant Response: Approximately 380 physical/virtual servers, 1,450 corporate and in-store workstations, and 620 point-of-sale terminals across all locations, plus roughly 900 handheld inventory scanners used in warehouse and back-of-store operations.',
])

paras = d.paragraphs
incident_idx = None
for i, p in enumerate(paras):
    if p.text.startswith('Applicant Response: Yes. In 2025'):
        incident_idx = i
        break
anchor = paras[incident_idx]
insert_after(anchor, d, heading='Prior Insurance & Claims History', level=2, body_paras=[
    'Who is your current/expiring cyber liability carrier, and what were the expiring limits and retention?',
    'Applicant Response: Beazley Insurance Company, expiring $15,000,000 aggregate limit, claims-made form, with a $100,000 per-claim retention. Policy period September 1, 2025 - September 1, 2026. We are moving away from Beazley because of the premium increase they proposed at this renewal following the ValueStore ransomware event.',
    'Has any insurer ever cancelled, non-renewed, or materially restricted your cyber coverage?',
    'Applicant Response: No. Beazley has offered to renew; we are voluntarily marketing the account for pricing and coverage reasons, not because coverage was declined or restricted.',
    'Please list all cyber-related claims, regardless of whether a loss run entry exists, over the past 5 years.',
    'Applicant Response: Two: (1) a January 2023 incident involving an encrypted company laptop lost by a regional sales manager (no data exfiltration confirmed; reported out of an abundance of caution), and (2) the June 2025 ValueStore legacy server ransomware event described above. Both are reflected in the attached Loss Run History.',
])

paras = d.paragraphs
compliance_idx = None
for i, p in enumerate(paras):
    if p.text.startswith('Applicant Response: No. We strictly operate'):
        compliance_idx = i
        break
anchor = paras[compliance_idx]
last = insert_after(anchor, d, heading='Vendor & Supply Chain Risk Management', level=2, body_paras=[
    'Do any third parties (contractors, POS vendors, supply chain partners) have remote access to your network?',
    'Applicant Response: Yes. Our HVAC/refrigeration maintenance contractors, our point-of-sale hardware vendor, and two legacy supply-chain EDI partners inherited from the ValueStore acquisition connect remotely via our Fortinet VPN to perform maintenance and data exchange. See the Ransomware Supplemental Application for details on how this access is authenticated.',
    'Do you require signed data security addenda or minimum security standards from vendors with network access?',
    'Applicant Response: Yes, for all vendors onboarded since 2023. We are in the process of retrofitting the same security addenda onto the legacy ValueStore vendor contracts, which we expect to complete by mid-2027.',
])

insert_after(last, d, heading='Business Continuity & Disaster Recovery', level=2, body_paras=[
    'Do you maintain a documented and tested Business Continuity / Disaster Recovery (BCDR) plan?',
    'Applicant Response: Yes. Our BCDR plan was last formally tested via tabletop exercise in March 2026 and covers store-level POS failover, corporate ERP recovery, and e-commerce site failover to a secondary AWS region. Recovery Time Objective (RTO) for the primary ERP system is 8 hours; Recovery Point Objective (RPO) is 1 hour.',
    'Do you carry a separate Property or Business Interruption policy that could respond to a cyber-triggered outage?',
    'Applicant Response: Yes, our commercial property policy contains a limited (sub-$250,000) cyber-triggered contingent business interruption endorsement, but our primary expectation is that first-party BI arising from a cyber event will be addressed under this cyber policy.',
])

# Officer certification at the very end
last_para = d.paragraphs[-1]
cert_heading = d.add_paragraph('Applicant Certification', style='Heading 2')
last_para._p.addnext(cert_heading._p)
cert_body = d.add_paragraph(
    'The undersigned officer certifies that the information contained in this application, and all attachments and supplements '
    'submitted herewith, is true and complete to the best of their knowledge and belief, and that no material facts have been '
    'omitted. The applicant agrees to notify the underwriter promptly of any material change in the information provided prior '
    'to the binding of coverage.'
)
cert_heading._p.addnext(cert_body._p)
sig_block = d.add_paragraph('Signed: Michael Reyes, Chief Financial Officer, NovaMart Inc.\nDate: August 22, 2026')
cert_body._p.addnext(sig_block._p)

cp = d.core_properties
cp.author = 'Sarah Thompson'
cp.last_modified_by = 'Sarah Thompson'

d.save(path)

d2 = Document(path)
print("saved, total paragraphs:", len(d2.paragraphs))
total_words = sum(len(p.text.split()) for p in d2.paragraphs)
print("total words:", total_words)
