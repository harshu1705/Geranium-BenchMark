import sys
from docx import Document
import re

def verify():
    doc = Document('Underwriting_Memo_NovaMart.docx')
    text = ' '.join([p.text for p in doc.paragraphs])
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text += ' ' + cell.text
    
    checks = {
        '$502,400,000': False,
        '$25,120': False,
        '$657,500': False,
        '$37,680': False,
        '1.20x': False,
        '1.25x': False,
        'CVE-2023-4528': False,
        'Section 8': False,
        'immediate file closure': False,
        'multiple simultaneous decline conditions': False,
        'no VP referral': False
    }

    for key in checks.keys():
        if key in text:
            checks[key] = True
        else:
            print(f"FAILED: Missing '{key}'")

    print("\n--- MATRIX ---")
    for k, v in checks.items():
        print(f"{k}: {'PASS' if v else 'FAIL'}")
    
    # check for truncated numbers
    truncated = re.findall(r'(\s,\d{3}(?:,\d{3})*)', text)
    if truncated:
        print(f"FAILED: Found truncated numbers: {truncated}")
    else:
        print("No truncated numbers found.")

if __name__ == '__main__':
    verify()
