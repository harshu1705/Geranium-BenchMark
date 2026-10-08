import os
import openpyxl
from docx import Document

def dump_docx(filepath, out_f):
    out_f.write(f"\n{'='*50}\nFILE: {filepath}\n{'='*50}\n")
    try:
        doc = Document(filepath)
        for p in doc.paragraphs:
            out_f.write(p.text + '\n')
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    row_text.append(cell.text.replace('\n', ' '))
                out_f.write(' | '.join(row_text) + '\n')
    except Exception as e:
        out_f.write(f"Error reading {filepath}: {e}\n")

def dump_xlsx(filepath, out_f):
    out_f.write(f"\n{'='*50}\nFILE: {filepath}\n{'='*50}\n")
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        for ws in wb.worksheets:
            out_f.write(f"\n--- SHEET: {ws.title} ---\n")
            for row in ws.iter_rows(values_only=True):
                out_f.write(str([str(cell) if cell is not None else '' for cell in row]) + '\n')
    except Exception as e:
        out_f.write(f"Error reading {filepath}: {e}\n")

with open('all_contents.md', 'w', encoding='utf-8') as out_f:
    input_dir = r"C:\Users\user\OneDrive\Desktop\geranium project\NovaMart_extracted\Inputs"
    for file in os.listdir(input_dir):
        filepath = os.path.join(input_dir, file)
        if file.endswith('.docx'):
            dump_docx(filepath, out_f)
        elif file.endswith('.xlsx'):
            dump_xlsx(filepath, out_f)

    golden_dir = r"C:\Users\user\OneDrive\Desktop\geranium project\NovaMart_extracted\Golden"
    for file in os.listdir(golden_dir):
        filepath = os.path.join(golden_dir, file)
        if file.endswith('.docx'):
            dump_docx(filepath, out_f)
